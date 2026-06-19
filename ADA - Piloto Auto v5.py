import os
import sys
import pandas as pd
import pyodbc
import re
import tkinter as tk
from tkinter import filedialog
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from ttkbootstrap.dialogs import Messagebox
from ttkbootstrap.widgets import DateEntry
from datetime import datetime, timedelta
import time
import subprocess
import locale
import logging
import threading
import math
import json
import glob
from PIL import Image, ImageTk
import smtplib
import socket
from email.message import EmailMessage

# --- CONFIGURACIÓN ---
logging.basicConfig(filename='app.log', level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
AUTOR = "© 2025 Víctor Domínguez. Todos los derechos reservados."
VERSION = "5.0.0"
password = 'oNer00FooR3n0'
SETTINGS_FILE = 'settings.json'
# Correo del administrador que recibirá las notificaciones de error.
ADMIN_EMAIL = "vdominguez@infoplazas.org.pa"

# Asocia el nombre de la regional con el nombre real de su carpeta
REGION_FOLDER_MAP = {
    "Los Santos": "BD_Azuero",
    "Chiriquí": "BD_Chiriqui",
    "Veraguas": "BD_Veraguas",
    "Panamá": "BD_Panamá"
}

import gspread
from google.oauth2.service_account import Credentials  # <-- AÑADE ESTA LÍNEA
from googleapiclient.discovery import build
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
import sucursales_history # <-- Módulo Nuevo Historial

try:
    locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')
except locale.Error:
    try:
        locale.setlocale(locale.LC_TIME, 'Spanish_Spain.1252')
    except:
        logging.warning("No se pudo establecer el locale en español.")


def asegurar_archivo_local(archivo, app_instance=None, tiempo_maximo=60, espera=5):
    """
    Fuerza la descarga de un archivo de OneDrive si está solo en la nube.
    Usa attrib de Windows para 'pin' el archivo localmente y mantiene la animación activa.
    
    tiempo_maximo: segundos totales a esperar antes de rendirse.
    espera: segundos entre cada intento.
    """
    tiempo_inicio = time.time()
    intento = 0
    
    while time.time() - tiempo_inicio < tiempo_maximo:
        intento += 1
        try:
            with open(archivo, 'rb') as f:
                f.read(1)
            return True  # ✅ ya está local
        except OSError:
            try:
                subprocess.run(['attrib', '+p', archivo], check=True, shell=True)
                logging.info(f"Forzando descarga de OneDrive para: {archivo}")
            except Exception as e:
                logging.error(f"No se pudo aplicar attrib a {archivo}: {e}")

            # 🔄 Mantener la animación viva durante la espera
            if app_instance and app_instance.animation_window and app_instance.animation_window.winfo_exists():
                app_instance.animation_window.update_status(
                    f"📥 Descargando desde OneDrive... intento {intento} (esperando {espera}s)"
                )
                app_instance.update_idletasks()

            time.sleep(espera)

    return False  # ❌ se alcanzó el tiempo máximo sin éxito


def _calcular_meses_relativos(dias):
    """Convierte un número de días en un texto relativo de meses."""
    if not isinstance(dias, (int, float)) or dias <= 30:
        return ""
    meses = dias // 30
    if meses == 1:
        return " (más de 1 mes)"
    else:
        return f" (más de {meses} meses)"



# --- LÓGICA DE NEGOCIO Y UTILIDADES ---

# --- LIBRERÍA DE SERVICIOS CON PRIORIDAD Y VARIANTES ---
LISTA_SERVICIOS_PRIORIZADA = [
    {'categoria': 'TALLER', 'variantes': ['TALLER','CURSO', 'CURSOS', 'CAPACITACION','CAPACITACIÓN','CAPACITACIONES']},
    {'categoria': 'REUNIÓN', 'variantes': ['REUNION', 'REUNIÓN','REUNIONES', 'MEETING']},
    {'categoria': 'CONSULTA', 'variantes': ['CONSULTA', 'CONSULTAS', 'ASESORIA', 'ASESORÍA']},
    {'categoria': 'VENTA', 'variantes': ['VENTA', 'VENTAS', 'PRODUCTO']},
    {'categoria': 'SCAN', 'variantes': ['SCAN', 'ESCANER', 'ESCANEADO','DIGITALIZACIÓN', 'SCANEOS']},
    {'categoria': 'CORREO', 'variantes': ['CORREO', 'EMAIL', 'E-MAIL', 'CORREOS']},
    {'categoria': 'TEL', 'variantes': ['TEL', 'TRAMITES EN LINEA','RECORD POLICIVO', 'TRAMITE EN LINEA']},
    {'categoria': 'LT', 'variantes': ['LT', 'LEVANTAMIENTO', 'LIBRE TECNOLOGÍA']},
    {'categoria': 'IMPRESIÓN', 'variantes': ['IMPRESION', 'IMPRESIÓN', 'IMPRESIONES', 'PRINT']},
    {'categoria': 'COPIA', 'variantes': ['COPIA', 'COPIAS', 'FOTOCOPIA', 'FOTOCOPIAS']},
    {'categoria': 'CINE', 'variantes': ['CINE', 'DOCUMENTAL', 'PELICULAS', 'PELÍCULA', 'CINE FORO','FORO', 'FOROS']}
]

SERVICIOS = [item['categoria'] for item in LISTA_SERVICIOS_PRIORIZADA]

def extraer_servicio(itemname, useraccount_dict=None):
    if useraccount_dict is None:
        useraccount_dict = {}
        
    itemname_upper = str(itemname).upper()
    palabras_item = itemname_upper.replace('-', ' ').split()

    item_key = str(itemname).strip().upper()
    if item_key in useraccount_dict:
        return 'USO DE PC'
        
    patrones_pc = [r'^PC\d*$', r'^LAPTOP\d*$', r'^CRON[OÓ]METRO\d*$', r'^HORAS?$']
    for palabra in palabras_item:
        for patron in patrones_pc:
            if re.match(patron, palabra, re.IGNORECASE):
                return 'USO DE PC'

    for servicio_info in LISTA_SERVICIOS_PRIORIZADA:
        categoria_actual = servicio_info['categoria']
        
        for variante in servicio_info['variantes']:
            # Regla Estricta: Si el servicio es corto (como LT o TEL), busca coincidencia exacta.
            if categoria_actual in ['LT', 'TEL', 'SCAN']: 
                if any(palabra == variante for palabra in palabras_item):
                    return categoria_actual
            # Regla Flexible: Para los demás servicios, usa startswith para admitir plurales.
            else:
                if any(palabra.startswith(variante) for palabra in palabras_item):
                    return categoria_actual
            
    return 'OTROS'

def generar_resumen_servicios(datos):
    if datos.empty or 'SERVICIO' not in datos.columns: 
        return pd.DataFrame()
    
    # Hacer una copia para no alterar el DataFrame original
    df_copy = datos.copy()
    
    # Asegurar que existan columnas de fecha y que sean de tipo datetime
    df_copy['year'] = df_copy['DATETIME'].dt.year
    df_copy['month'] = df_copy['DATETIME'].dt.month
    
    # Agrupar y pivotar por sucursal y mes
    pivot_servicios = df_copy.pivot_table(
        index=['source_folder', 'year', 'month'], 
        columns='SERVICIO', 
        aggfunc='size', 
        fill_value=0
    ).reset_index()
    
    # Asegurar que todas las columnas de categorías estén presentes
    for servicio in SERVICIOS + ['OTROS', 'USO DE PC']:
        if servicio not in pivot_servicios.columns:
            pivot_servicios[servicio] = 0
            
    # Agregar la columna de mes formateado en texto (español)
    pivot_servicios['Mes'] = pivot_servicios['month'].apply(lambda x: datetime(1900, x, 1).strftime('%B').capitalize())
    
    # Renombrar source_folder a Infoplaza y year a Año
    pivot_servicios.rename(columns={'source_folder': 'Infoplaza', 'year': 'Año'}, inplace=True)
    
    # Calcular el total de servicios por fila
    columnas_servicios = SERVICIOS + ['OTROS', 'USO DE PC']
    pivot_servicios['Total'] = pivot_servicios[columnas_servicios].sum(axis=1)
    
    # Ordenar y seleccionar columnas finales
    columnas_finales = ['Infoplaza', 'Año', 'Mes'] + columnas_servicios + ['Total']
    return pivot_servicios[columnas_finales]

def resource_path(relative_path):
    try: base_path = sys._MEIPASS
    except Exception: base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

# --- LIMPIAR DE CARACTERES INVÁLIDOS O NO LEGIBLES
def limpiar_caracteres_invalidos(df):
    """
    Recorre un DataFrame y elimina los caracteres de control inválidos para XML/Excel
    de todas las columnas de tipo string.
    """
    if df.empty:
        return df

    # Expresión regular para encontrar caracteres de control inválidos
    # (excluyendo tab, salto de línea y retorno de carro, que son válidos en Excel)
    regex_control_chars = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1F]')

    def limpiar_celda(valor):
        if isinstance(valor, str):
            return regex_control_chars.sub('', valor)
        return valor

    # Aplica la limpieza a cada columna de tipo 'object' (generalmente strings)
    for col in df.select_dtypes(include=['object']).columns:
        df[col] = df[col].astype(str).apply(limpiar_celda)
    
    return df

def obtener_regional(nombre_carpeta_raiz):
    """ Determina la regional basada en el nombre de la carpeta raíz. """
    if "Azuero" in nombre_carpeta_raiz: return "Los Santos"
    if "Chiriqui" in nombre_carpeta_raiz or "Chiriquí" in nombre_carpeta_raiz: return "Chiriquí"
    if "Veraguas" in nombre_carpeta_raiz: return "Veraguas"
    if "Panamá" in nombre_carpeta_raiz: return "Panamá"
    return "Desconocida"


def extraer_datos_db(db_path, password, start_date, end_date, app_instance):
    if app_instance and app_instance.animation_window:
        nombre_carpeta = os.path.basename(os.path.dirname(db_path))
        app_instance.animation_window.update_status(f"Analizando: {nombre_carpeta}")
    conn_str = f'DRIVER={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};PWD={password};'
    try:
        with pyodbc.connect(conn_str, autocommit=True) as conn:
            query = "SELECT DATETIME, USERNAME, ITEMNAME FROM SALES WHERE [DATETIME] >= ? AND [DATETIME] < ?"
            data = pd.read_sql(query, conn, params=[start_date, end_date])
            if not data.empty: data['DATETIME'] = pd.to_datetime(data['DATETIME'])
            
            cur = conn.cursor()
            cur.execute("SELECT USERNAME, SEX FROM USERACCOUNT")
            useraccount_dict = {}
            for row in cur.fetchall():
                if row.USERNAME and row.SEX:
                    useraccount_dict[str(row.USERNAME).strip().upper()] = str(row.SEX).strip().upper()
                    
            return data, useraccount_dict, None
    except Exception as e:
        logging.error(f"Error extrayendo datos de {db_path}: {e}"); return None, {}, str(e)

def analizar_itemname(data, useraccount_dict=None):
    if useraccount_dict is None:
        useraccount_dict = {}
        
    alternador_sexo = True
    def limpiar_itemname(itemname): return ' '.join(str(itemname).replace('--', '-').strip('-').strip().split())
    
    def obtener_sexo(itemname):
        nonlocal alternador_sexo
        item_key = str(itemname).strip().upper()
        if item_key in useraccount_dict:
            sexo_db = useraccount_dict[item_key]
            if sexo_db in ['M', 'F']:
                return sexo_db
                
        itemname_limpio = limpiar_itemname(itemname).upper(); partes = [p.strip() for p in itemname_limpio.split('-')]
        if partes and partes[-1] in ['M', 'F']: return partes[-1]
        if len(partes) >= 2 and partes[-2] in ['M', 'F']: return partes[-2]
        alternador_sexo = not alternador_sexo
        return 'M' if alternador_sexo else 'F'
        
    def obtener_tipo_u(itemname):
        itemname_limpio = limpiar_itemname(itemname).upper(); partes = [p.strip() for p in itemname_limpio.split('-')]
        tipos_validos = ['P', 'S', 'U', 'D', 'TE', 'PG']
        if len(partes) >= 2 and partes[-2] in tipos_validos: return partes[-2]
        if len(partes) >= 3 and partes[-3] in tipos_validos: return partes[-3]
        return 'PG'
        
    def obtener_observacion(itemname):
        item_key = str(itemname).strip().upper()
        if item_key in useraccount_dict and useraccount_dict[item_key] in ['M', 'F']:
            return 'CORREGIDO VIA DB'
            
        itemname_limpio = limpiar_itemname(itemname).upper(); partes = [p.strip() for p in itemname_limpio.split('-')]
        tipos_validos = ['P', 'S', 'U', 'D', 'TE', 'PG']
        if len(partes) >= 2 and partes[-1] in ['M', 'F'] and partes[-2] in tipos_validos: return 'OK'
        return 'VERIFICAR'
        
    data['SEXO'] = data['ITEMNAME'].apply(obtener_sexo); data['TIPO U'] = data['ITEMNAME'].apply(obtener_tipo_u); data['OBSERVACIÓN'] = data['ITEMNAME'].apply(obtener_observacion)
    data['SERVICIO'] = data['ITEMNAME'].apply(lambda x: extraer_servicio(x, useraccount_dict))
    return data

def procesar_bases_de_datos(ruta_carpeta_raiz, password, start_date, end_date, app_instance):
    """
    Procesa todas las subcarpetas asegurando que:
    - Se use SIEMPRE la FECHA DE MODIFICACIÓN del archivo como referencia.
    - Solo se descargue y espere el archivo .mdb más reciente (OneDrive).
    - Se mantenga un reporte de incidentes detallado.
    - Se rellenen los meses faltantes SOLO dentro del rango [start_date, end_date].
    - Se limpien caracteres inválidos antes de retornar.
    """

    subcarpetas = [f.path for f in os.scandir(ruta_carpeta_raiz) if f.is_dir()]
    todos_los_datos, datos_incidentes = [], []
    fecha_sincronizacion_actual = datetime.now().date()

    for i, carpeta in enumerate(subcarpetas):
        # Estado inicial en interfaz
        if app_instance and app_instance.animation_window and app_instance.animation_window.winfo_exists():
            app_instance.animation_window.update_status(
                f"🔄 Procesando: {os.path.basename(carpeta)} ({i+1}/{len(subcarpetas)})"
            )

        archivos_mdb = [f for f in os.listdir(carpeta) if f.endswith('.mdb')]
        archivo_seleccionado, fecha_archivo_dt = None, None

        if archivos_mdb:
            archivos_con_fechas_modificacion = [
                (os.path.join(carpeta, a), datetime.fromtimestamp(os.path.getmtime(os.path.join(carpeta, a))))
                for a in archivos_mdb
            ]
            archivos_con_fechas_modificacion.sort(key=lambda x: x[1], reverse=True)
            archivo_seleccionado = archivos_con_fechas_modificacion[0][0]
            fecha_archivo_dt = archivos_con_fechas_modificacion[0][1].date()
            fecha_archivo_str = fecha_archivo_dt.strftime("%d/%m/%Y")
            dias_sin_sinc = (fecha_sincronizacion_actual - fecha_archivo_dt).days
            if dias_sin_sinc < 0:  # Seguridad contra fechas futuras
                dias_sin_sinc = 0
        else:
            fecha_archivo_str, dias_sin_sinc = "N/A", "N/A"

        # Carpeta cerrada definitivamente
        archivos_cerrada_def = glob.glob(os.path.join(carpeta, "CERRADA DEFINITIVAMENTE.txt"))
        if archivos_cerrada_def:
            obs = os.path.splitext(os.path.basename(archivos_cerrada_def[0]))[0]
            datos_incidentes.append([datetime.now(), os.path.basename(carpeta), fecha_archivo_str,
                                     dias_sin_sinc, "N/A", "N/A", obs])
            continue

        # Detectar si hay otros archivos .txt (observaciones)
        # NUEVA LÓGICA: Guardar observación pero NO agregar registro todavía
        observacion_txt = None
        otros_archivos_cerrada = glob.glob(os.path.join(carpeta, "*.txt"))
        if otros_archivos_cerrada:
            observacion_txt = os.path.splitext(os.path.basename(otros_archivos_cerrada[0]))[0]
            # NO agregar registro aquí - se agregará al procesar el .mdb

        # Carpeta sin MDB
        if not archivos_mdb:
            # Si hay observación de .txt, usarla; si no, "Carpeta Vacía"
            obs_final = observacion_txt if observacion_txt else "Carpeta Vacía"
            datos_incidentes.append([datetime.now(), os.path.basename(carpeta), "N/A",
                                     "N/A", "N/A", "N/A", obs_final])
            continue

        # 🔒 BLOQUE ONEDRIVE: asegura que el archivo esté descargado
        if not asegurar_archivo_local(archivo_seleccionado, app_instance, tiempo_maximo=60, espera=5):
            logging.error(f"No se pudo descargar el archivo desde OneDrive: {archivo_seleccionado}")
            datos_incidentes.append([datetime.now(), os.path.basename(carpeta), fecha_archivo_str,
                                     dias_sin_sinc, "N/A", "N/A", "⚠️ Error OneDrive"])
            continue

        # ✨ Mensaje de transición → archivo ya está listo
        if app_instance and app_instance.animation_window and app_instance.animation_window.winfo_exists():
            app_instance.animation_window.update_status("✅ Archivo listo, procesando datos...")
            app_instance.update_idletasks()

        # Procesamiento MDB
        datos, useraccount_dict, error = extraer_datos_db(archivo_seleccionado, password, start_date, end_date, app_instance)

        if datos is not None and not datos.empty:
            # Filtrar registros con fechas anteriores a 1900
            if 'DATETIME' in datos.columns and pd.api.types.is_datetime64_any_dtype(datos['DATETIME']):
                datos = datos[datos['DATETIME'].dt.year >= 1900]

        if datos is not None and not datos.empty:
            datos['source_folder'] = os.path.basename(carpeta)
            datos_analizados = analizar_itemname(datos, useraccount_dict)
            todos_los_datos.append(datos_analizados)
            
            # NUEVA LÓGICA: Si hay observación de .txt, usarla; si no, "OK"
            obs_final = observacion_txt if observacion_txt else "OK"
            
            datos_incidentes.append([
                datetime.now(), os.path.basename(carpeta), fecha_archivo_str, dias_sin_sinc,
                datos['DATETIME'].min().strftime("%d/%m/%Y"),
                datos['DATETIME'].max().strftime("%d/%m/%Y"),
                obs_final  # Usar observación del .txt si existe, sino "OK"
            ])
        else:
            # Si no hay datos, usar observación del .txt si existe, sino el error/sin registros
            if observacion_txt:
                obs_final = observacion_txt
            else:
                obs_final = f"Error: {error}" if error else "Sin registros para el periodo"
            
            datos_incidentes.append([datetime.now(), os.path.basename(carpeta), fecha_archivo_str,
                                     dias_sin_sinc, "N/A", "N/A", obs_final])

    # Construir reporte incidentes
    reporte_incidentes_df = pd.DataFrame(datos_incidentes, columns=[
        'Fecha y Hora', 'Carpeta', 'Fecha Archivo', 'Días sin Sincronizar',
        'Primer Registro', 'Último Registro', 'Observación'
    ])

    if not todos_los_datos:
        return pd.DataFrame(), reporte_incidentes_df

    # Combinar datos
    datos_combinados = pd.concat(todos_los_datos, ignore_index=True)

    # Relleno de meses faltantes SOLO dentro del rango
    end_date_para_rango = end_date - timedelta(days=1)
    rango_meses = pd.date_range(start=start_date, end=end_date_para_rango, freq='MS').strftime("%Y-%m").tolist()

    if 'DATETIME' in datos_combinados.columns:
        datos_combinados['MesAnio'] = datos_combinados['DATETIME'].dt.strftime('%Y-%m')

    infoplazas = datos_combinados['source_folder'].unique()
    datos_faltantes = []

    for infoplaza in infoplazas:
        datos_infoplaza = datos_combinados[datos_combinados['source_folder'] == infoplaza]
        meses_existentes = datos_infoplaza['MesAnio'].unique() if 'MesAnio' in datos_infoplaza.columns else []
        for mes_anio in rango_meses:
            if mes_anio not in meses_existentes:
                anio, mes = map(int, mes_anio.split('-'))
                datos_faltantes.append({
                    'DATETIME': datetime(year=anio, month=mes, day=1),
                    'USERNAME': 'Sin Registro', 'ITEMNAME': 'Sin Registro',
                    'SEXO': 'S/D', 'TIPO U': 'S/D', 'OBSERVACIÓN': 'MES SIN DATOS',
                    'SERVICIO': 'S/D',
                    'source_folder': infoplaza, 'MesAnio': mes_anio
                })

    if datos_faltantes:
        df_faltantes = pd.DataFrame(datos_faltantes)
        datos_combinados = pd.concat([datos_combinados, df_faltantes], ignore_index=True)

    # Ordenar datos finales
    datos_combinados = datos_combinados.sort_values(by=['source_folder', 'DATETIME']).reset_index(drop=True)

    if 'MesAnio' in datos_combinados.columns:
        datos_combinados = datos_combinados.drop(columns=['MesAnio'])

    # Limpieza final
    datos_combinados = limpiar_caracteres_invalidos(datos_combinados)
    reporte_incidentes_df = limpiar_caracteres_invalidos(reporte_incidentes_df)

    return datos_combinados, reporte_incidentes_df
        
   

def generar_resumen(datos):
    if datos.empty: return pd.DataFrame()
    datos['year'] = datos['DATETIME'].dt.year; datos['month'] = datos['DATETIME'].dt.month
    resumen_agg = datos.groupby(['source_folder', 'year', 'month']).agg(M=('SEXO', lambda x: (x == 'M').sum()), F=('SEXO', lambda x: (x == 'F').sum()), P=('TIPO U', lambda x: (x == 'P').sum()), S=('TIPO U', lambda x: (x == 'S').sum()), U=('TIPO U', lambda x: (x == 'U').sum()), D=('TIPO U', lambda x: (x == 'D').sum()), TE=('TIPO U', lambda x: (x == 'TE').sum()), PG=('TIPO U', lambda x: (x == 'PG').sum())).reset_index()
    resumen_agg['Mes'] = resumen_agg['month'].apply(lambda x: datetime(1900, x, 1).strftime('%B').capitalize())
    resumen_agg.rename(columns={'source_folder': 'Infoplaza', 'year': 'Año', 'M': 'Masculino', 'F': 'Femenino', 'P': 'Primaria', 'S': 'Secundaria', 'U': 'Universitario', 'D': 'Docente', 'TE': 'Tercera Edad', 'PG': 'Público General'}, inplace=True)
    resumen_agg['Total'] = resumen_agg[['Primaria', 'Secundaria', 'Universitario', 'Docente', 'Tercera Edad', 'Público General']].sum(axis=1)
    return resumen_agg[['Infoplaza', 'Año', 'Mes', 'Masculino', 'Femenino', 'Primaria', 'Secundaria', 'Universitario', 'Docente', 'Tercera Edad', 'Público General', 'Total']]

def generar_resumen_cruzado(datos):
    if datos.empty: return pd.DataFrame()
    
    # Crear un DataFrame temporal para no afectar el original
    df = datos.copy()
    if 'year' not in df.columns:
        df['year'] = df['DATETIME'].dt.year
    if 'month' not in df.columns:
        df['month'] = df['DATETIME'].dt.month
        
    # Filtrar meses sin datos
    df = df[df['TIPO U'] != 'S/D']
    if df.empty: return pd.DataFrame()

    # Agrupar y calcular totales
    resumen_agg = df.groupby(['source_folder', 'year', 'month', 'TIPO U']).agg(
        Masculino=('SEXO', lambda x: (x == 'M').sum()),
        Femenino=('SEXO', lambda x: (x == 'F').sum())
    ).reset_index()

    resumen_agg['Total'] = resumen_agg['Masculino'] + resumen_agg['Femenino']
    resumen_agg['Mes'] = resumen_agg['month'].apply(lambda x: datetime(1900, x, 1).strftime('%B').capitalize())

    # Mapear TIPO U
    tipo_map = {
        'P': 'Primaria', 'S': 'Secundaria', 'U': 'Universitario',
        'D': 'Docente', 'TE': 'Tercera Edad', 'PG': 'Público General'
    }
    resumen_agg['TIPO U'] = resumen_agg['TIPO U'].map(tipo_map).fillna(resumen_agg['TIPO U'])

    resumen_agg.rename(columns={'source_folder': 'Infoplaza', 'year': 'Año', 'TIPO U': 'Tipo Usuario'}, inplace=True)
    
    return resumen_agg[['Infoplaza', 'Año', 'Mes', 'Tipo Usuario', 'Masculino', 'Femenino', 'Total']]

def guardar_en_excel(datos, reporte_incidentes, resumen, archivo_salida, resumen_cruzado=None):
    import traceback
    try:
        resumen_servicios = generar_resumen_servicios(datos)
        with pd.ExcelWriter(archivo_salida, engine='openpyxl') as writer:
            # Dividir "Datos Completos" si excede el límite de filas de Excel (aprox 1M)
            limite_filas = 1000000
            if len(datos) > limite_filas:
                num_partes = (len(datos) // limite_filas) + 1
                for i in range(num_partes):
                    inicio = i * limite_filas
                    fin = inicio + limite_filas
                    chunk = datos.iloc[inicio:fin]
                    nombre_hoja = f'Datos Completos P{i+1}'
                    chunk.to_excel(writer, sheet_name=nombre_hoja, index=False)
            else:
                datos.to_excel(writer, sheet_name='Datos Completos', index=False)
                
            resumen.to_excel(writer, sheet_name='Resumen', index=False)
            if not resumen_servicios.empty:
                resumen_servicios.to_excel(writer, sheet_name='Servicios', index=False)
            if resumen_cruzado is not None and not resumen_cruzado.empty:
                resumen_cruzado.to_excel(writer, sheet_name='Resumen Tipo usuario y género', index=False)
                
                # Fase 4: Formateo básico a la nueva sheet
                worksheet = writer.sheets['Resumen Tipo usuario y género']
                from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
                from openpyxl.utils import get_column_letter
                header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                
                for col_idx, col in enumerate(resumen_cruzado.columns, 1):
                    cell = worksheet.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border
                    
                    max_length = max(resumen_cruzado[col].astype(str).map(len).max(), len(col)) + 2
                    worksheet.column_dimensions[get_column_letter(col_idx)].width = max_length
                
                for r_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=7), 2):
                    for cell in row:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0'

            reporte_incidentes.to_excel(writer, sheet_name='Reporte Incidentes', index=False)
        logging.info(f"Archivo guardado en {archivo_salida}"); return True
    except Exception as e: 
        logging.error(f"Error al guardar Excel: {e}\n{traceback.format_exc()}")
        return False

def redimensionar_hoja_google_sheets(client, sheet_id, sheet_name, num_rows, num_cols):
    try:
        sheet_metadata = client.spreadsheets().get(spreadsheetId=sheet_id).execute()
        sheets = sheet_metadata.get('sheets', '')
        sheet_id_num = None
        for s in sheets:
            if s['properties']['title'] == sheet_name:
                sheet_id_num = s['properties']['sheetId']
                break
        if sheet_id_num is not None:
            num_rows = max(1, num_rows)
            num_cols = max(1, num_cols)
            body = {
                'requests': [
                    {
                        'updateSheetProperties': {
                            'properties': {
                                'sheetId': sheet_id_num,
                                'gridProperties': {
                                    'rowCount': num_rows,
                                    'columnCount': num_cols
                                }
                            },
                            'fields': 'gridProperties(rowCount,columnCount)'
                        }
                    }
                ]
            }
            client.spreadsheets().batchUpdate(spreadsheetId=sheet_id, body=body).execute()
    except Exception as re:
        logging.warning(f"No se pudo redimensionar la hoja '{sheet_name}': {re}")

def subir_datos_a_google_sheets(resumen_df, nombre_carpeta_raiz):
    try:
        creds_path = resource_path('credentials_gebilo.json');
        if not os.path.exists(creds_path): raise FileNotFoundError(f"No se encontró el archivo de credenciales: {creds_path}")
        scope = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]; creds = Credentials.from_service_account_file(creds_path, scopes=scope)
        client = build('sheets', 'v4', credentials=creds); sheet_id = '12GPBc0JAMRY2b7YMWcFDGiAP9xOrfDNgoZdiwQyuWdc'
        sheet_metadata = client.spreadsheets().get(spreadsheetId=sheet_id).execute(); sheets = sheet_metadata.get('sheets', '')
        if not any(s['properties']['title'] == nombre_carpeta_raiz for s in sheets):
            body = {'requests': [{'addSheet': {'properties': {'title': nombre_carpeta_raiz}}}]}; client.spreadsheets().batchUpdate(spreadsheetId=sheet_id, body=body).execute()
        range_name = f"{nombre_carpeta_raiz}!A:Z"; result = client.spreadsheets().values().get(spreadsheetId=sheet_id, range=range_name).execute(); values = result.get('values', [])
        numeric_cols = ['Año', 'Masculino', 'Femenino', 'Primaria', 'Secundaria', 'Universitario', 'Docente', 'Tercera Edad', 'Público General', 'Total']
        if values and len(values[0]) > 0 and 'Infoplaza' in values[0]:
            df_hoja = pd.DataFrame(values[1:], columns=values[0])
            for col in numeric_cols:
                if col in df_hoja.columns: df_hoja[col] = pd.to_numeric(df_hoja[col], errors='coerce').fillna(0).astype(int)
        else: df_hoja = pd.DataFrame(columns=resumen_df.columns)
        for _, fila_resumen in resumen_df.iterrows():
            fila_resumen['Año'] = int(fila_resumen['Año']); fila_resumen['Total'] = int(fila_resumen['Total'])
            condicion = ((df_hoja['Infoplaza'].astype(str) == str(fila_resumen['Infoplaza'])) & (df_hoja['Año'] == fila_resumen['Año']) & (df_hoja['Mes'].astype(str) == str(fila_resumen['Mes'])))
            coincidencias = df_hoja[condicion]
            if coincidencias.empty: df_hoja = pd.concat([df_hoja, pd.DataFrame([fila_resumen])], ignore_index=True)
            else:
                indice = coincidencias.index[0]
                if df_hoja.loc[indice, 'Total'] < fila_resumen['Total']:
                    df_hoja = df_hoja.drop(indice); df_hoja = pd.concat([df_hoja, pd.DataFrame([fila_resumen])], ignore_index=True)
        columnas_a_sumar = ['Primaria', 'Secundaria', 'Universitario', 'Docente', 'Tercera Edad', 'Público General']
        for col in columnas_a_sumar:
            if col not in df_hoja.columns: df_hoja[col] = 0
            df_hoja[col] = pd.to_numeric(df_hoja[col], errors='coerce').fillna(0).astype(int)
        df_hoja['Total'] = df_hoja[columnas_a_sumar].sum(axis=1)
        for col in numeric_cols:
             if col in df_hoja.columns: df_hoja[col] = pd.to_numeric(df_hoja[col], errors='coerce').fillna(0).astype(int)
        valores_actualizados = [df_hoja.columns.tolist()] + df_hoja.where(pd.notna(df_hoja), None).values.tolist()
        
        # Redimensionar la hoja antes de escribir para liberar celdas del limite de 10M
        redimensionar_hoja_google_sheets(client, sheet_id, nombre_carpeta_raiz, len(valores_actualizados), len(valores_actualizados[0]))
        
        client.spreadsheets().values().clear(spreadsheetId=sheet_id, range=nombre_carpeta_raiz).execute()
        body = {'values': valores_actualizados}
        client.spreadsheets().values().update(spreadsheetId=sheet_id, range=f"{nombre_carpeta_raiz}!A1", valueInputOption='USER_ENTERED', body=body).execute()
        logging.info(f"Datos actualizados y subidos a la hoja: {nombre_carpeta_raiz}"); return True
    except Exception as e: logging.error(f"Error al subir a Google Sheets: {e}"); raise e

def subir_historial_sincronizacion(nuevo_historial_df, nombre_carpeta_raiz):
    try:
        # 1. Preparar el DataFrame para la subida
        df_a_subir = nuevo_historial_df.copy()
        df_a_subir['Fecha y Hora'] = pd.to_datetime(df_a_subir['Fecha y Hora']).dt.strftime('%Y-%m-%d %H:%M:%S')
        
        # Agregar columna temporal solo con la fecha (sin hora) para deduplicación
        df_a_subir['FechaAnalisis'] = pd.to_datetime(df_a_subir['Fecha y Hora']).dt.strftime('%Y-%m-%d')
        
        df_a_subir.rename(columns={
            'Carpeta': 'Sucursal', 
            'Fecha Archivo': 'FechaSincronizacion',
            'Días sin Sincronizar': 'DiasSinSinc'
        }, inplace=True)
        
        # Seleccionar columnas finales (sin incluir FechaAnalisis en Sheets)
        df_a_subir_final = df_a_subir[['Fecha y Hora', 'Sucursal', 'FechaSincronizacion', 'DiasSinSinc', 'Observación']]

        # 2. Autenticación y conexión
        creds_path = resource_path('credentials_gebilo.json')
        if not os.path.exists(creds_path): raise FileNotFoundError(f"No se encontró el archivo de credenciales: {creds_path}")
        scope = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_file(creds_path, scopes=scope)
        client = build('sheets', 'v4', credentials=creds)
        sheet_id = '12GPBc0JAMRY2b7YMWcFDGiAP9xOrfDNgoZdiwQyuWdc'
        sheet_name = f"Historial_{nombre_carpeta_raiz}"

        # 3. Asegurarse de que la pestaña (hoja) exista
        sheet_metadata = client.spreadsheets().get(spreadsheetId=sheet_id).execute(); sheets = sheet_metadata.get('sheets', '')
        if not any(s['properties']['title'] == sheet_name for s in sheets):
            body = {'requests': [{'addSheet': {'properties': {'title': sheet_name}}}]}; client.spreadsheets().batchUpdate(spreadsheetId=sheet_id, body=body).execute()

        # 4. Leer datos existentes y fusionar
        range_name = f"{sheet_name}!A:E"
        result = client.spreadsheets().values().get(spreadsheetId=sheet_id, range=range_name).execute(); values = result.get('values', [])
        
        df_google = pd.DataFrame()
        if values and len(values) > 1:
            header = values[0]
            num_cols = len(header)
            data_rows = [row + [None] * (num_cols - len(row)) for row in values[1:]]
            df_google = pd.DataFrame(data_rows, columns=header)
            # Agregar FechaAnalisis temporal al DataFrame existente
            if not df_google.empty and 'Fecha y Hora' in df_google.columns:
                df_google['FechaAnalisis'] = pd.to_datetime(df_google['Fecha y Hora'], errors='coerce').dt.strftime('%Y-%m-%d')

        # Combinar datos existentes con nuevos
        df_combined = pd.concat([df_google, df_a_subir], ignore_index=True)
        
        # NUEVA LÓGICA: Deduplicar por (Sucursal, FechaAnalisis) - mantener el último registro del día
        df_final = df_combined.drop_duplicates(subset=['Sucursal', 'FechaAnalisis'], keep='last')
        df_final = df_final.sort_values(by=['Sucursal', 'Fecha y Hora']).reset_index(drop=True)
        
        # Eliminar columna temporal antes de subir
        df_final = df_final[['Fecha y Hora', 'Sucursal', 'FechaSincronizacion', 'DiasSinSinc', 'Observación']]
        
        # 5. Escribir datos de vuelta a la hoja
        valores_actualizados = [df_final.columns.tolist()] + df_final.where(pd.notna(df_final), None).values.tolist()
        
        # Redimensionar la hoja antes de escribir para liberar celdas del limite de 10M
        redimensionar_hoja_google_sheets(client, sheet_id, sheet_name, len(valores_actualizados), len(valores_actualizados[0]))
        
        client.spreadsheets().values().clear(spreadsheetId=sheet_id, range=sheet_name).execute()
        body = {'values': valores_actualizados}
        client.spreadsheets().values().update(spreadsheetId=sheet_id, range=f"{sheet_name}!A1", valueInputOption='USER_ENTERED', body=body).execute()
        
        logging.info(f"Historial de sincronización subido a la hoja: {sheet_name}")
        return True
    except Exception as e:
        logging.error(f"Error al subir el historial de sincronización: {e}")
        return False


def subir_porcentajes_sincronizacion(reporte_df, regional):
    """
    Sube o actualiza los porcentajes de sincronización mensual a la hoja 'PorcSinc' en Google Sheets.
    
    Args:
        reporte_df (DataFrame): DataFrame con el reporte de incidentes
        regional (str): Nombre de la regional
    
    Returns:
        bool: True si se subió correctamente, False en caso de error
    """
    try:
        if reporte_df.empty:
            logging.warning("DataFrame de reporte vacío, no se subirán porcentajes.")
            return False
        
        # Fecha actual del reporte
        fecha_reporte = datetime.now()
        mes_reporte = fecha_reporte.month
        año_reporte = fecha_reporte.year
        fecha_reporte_str = fecha_reporte.strftime('%Y-%m-%d')
        
        # Separar cerradas de operativas (criterio usado para el correo)
        df_operativas = reporte_df[reporte_df['Observación'] != 'CERRADA DEFINITIVAMENTE'].copy()
        
        if df_operativas.empty:
            logging.warning("No hay sucursales operativas para calcular porcentajes.")
            return False
        
        # Asegurar columna numérica
        df_operativas['Días sin Sincronizar'] = pd.to_numeric(df_operativas['Días sin Sincronizar'], errors='coerce').fillna(0)
        
        # --- CLASIFICACIÓN EXACTA COMO EN EL CORREO ---
        
        # 1. Carpetas vacías
        df_carpetas_vacias = df_operativas[df_operativas['Observación'] == 'Carpeta Vacía'].copy()
        
        # 2. Preparar dataframe de trabajo (excluir carpetas vacías)
        df_trabajo = df_operativas[df_operativas['Observación'] != 'Carpeta Vacía'].copy()
        
        # 3. Clasificación secuencial
        observaciones_basicas = ['OK', 'Sin registros para el periodo', 'CERRADA DEFINITIVAMENTE', 'CERRADA TEMPORALMENTE']
        
        # 3a. Sucursales con observaciones relevantes (> 10 días + observación especial)
        df_con_observaciones = df_trabajo[
            (df_trabajo['Días sin Sincronizar'] > 10) & 
            (~df_trabajo['Observación'].isin(observaciones_basicas))
        ].copy()
        
        # 3b. Sucursales con > 10 días pero observación OK/Sin registros (excluir ya clasificadas)
        indices_clasificados = df_con_observaciones.index.tolist()
        df_revision_dias = df_trabajo[
            (~df_trabajo.index.isin(indices_clasificados)) & 
            (df_trabajo['Días sin Sincronizar'] > 10) & 
            (df_trabajo['Observación'].isin(['OK', 'Sin registros para el periodo']))
        ].copy()
        
        # 3c. Sucursales al día (<= 10 días)
        df_en_periodo = df_trabajo[
            (df_trabajo['Días sin Sincronizar'] <= 10)
        ].copy()
        
        # --- MÉTRICAS COMPATIBLES CON EL CORREO ---
        total_verificadas = len(df_operativas)
        
        # Al día: TODAS las operativas con <= 10 días (sin excluir carpetas vacías aquí)
        al_dia = len(df_operativas[df_operativas['Días sin Sincronizar'] <= 10])
        
        # Para revisión: Unión de Carpetas Vacías O Más de 10 días (para evitar doble conteo)
        df_para_revision = df_operativas[
            (df_operativas['Observación'] == 'Carpeta Vacía') | 
            (df_operativas['Días sin Sincronizar'] > 10)
        ]
        para_revision = len(df_para_revision)
        
        # Calcular porcentajes (Redondeado a 1 decimal para consistencia total con el correo)
        porc_al_dia = round((al_dia / total_verificadas) * 100, 1) if total_verificadas > 0 else 0
        porc_revision = round((para_revision / total_verificadas) * 100, 1) if total_verificadas > 0 else 0
        
        # Crear registro para PorcSinc
        nuevo_registro = {
            'Fecha': fecha_reporte_str,
            'Regional': regional,
            'Verificadas': total_verificadas,
            'Al día': al_dia,
            'Para revisión': para_revision,
            'Porc. al día': porc_al_dia,
            'Porc. Revisión': porc_revision
        }
        
        # 2. Autenticación y conexión a Google Sheets
        creds_path = resource_path('credentials_gebilo.json')
        if not os.path.exists(creds_path):
            raise FileNotFoundError(f"No se encontró el archivo de credenciales: {creds_path}")
        
        scope = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_file(creds_path, scopes=scope)
        client = build('sheets', 'v4', credentials=creds)
        sheet_id = '12GPBc0JAMRY2b7YMWcFDGiAP9xOrfDNgoZdiwQyuWdc'
        sheet_name = 'PorcSinc'
        
        # 3. Asegurarse de que la hoja PorcSinc exista
        sheet_metadata = client.spreadsheets().get(spreadsheetId=sheet_id).execute()
        sheets = sheet_metadata.get('sheets', '')
        
        if not any(s['properties']['title'] == sheet_name for s in sheets):
            # Crear la hoja si no existe
            body = {'requests': [{'addSheet': {'properties': {'title': sheet_name}}}]}
            client.spreadsheets().batchUpdate(spreadsheetId=sheet_id, body=body).execute()
            logging.info(f"Hoja '{sheet_name}' creada en Google Sheets.")
        
        # 4. Leer datos existentes
        range_name = f"{sheet_name}!A:G"
        result = client.spreadsheets().values().get(spreadsheetId=sheet_id, range=range_name).execute()
        values = result.get('values', [])
        
        # Crear DataFrame con datos existentes
        if values and len(values) > 1:
            header = values[0]
            num_cols = len(header)
            data_rows = [row + [None] * (num_cols - len(row)) for row in values[1:]]
            df_google = pd.DataFrame(data_rows, columns=header)
        else:
            # Si no hay datos, crear estructura inicial
            df_google = pd.DataFrame(columns=['Fecha', 'Regional', 'Verificadas', 'Al día', 'Para revisión', 'Porc. al día', 'Porc. Revisión'])
        
        # 5. Evitar duplicados el mismo día para la misma regional
        if not df_google.empty and 'Fecha' in df_google.columns:
            # Asegurar que 'Fecha' se trata como cadena para la comparación del día
            df_google['Fecha'] = df_google['Fecha'].astype(str)
            
            # Buscar coincidencia: misma regional y misma fecha exacta (día)
            condicion = (
                (df_google['Regional'].astype(str) == str(regional)) &
                (df_google['Fecha'] == fecha_reporte_str)
            )
            coincidencias = df_google[condicion]
            
            if not coincidencias.empty:
                # Actualizar el registro del mismo día
                indice = coincidencias.index[0]
                for key, value in nuevo_registro.items():
                    df_google.loc[indice, key] = value
                logging.info(f"Actualizado registro del día para {regional} - {fecha_reporte_str}")
            else:
                # Agregar nuevo registro (día diferente o regional diferente)
                nuevo_df = pd.DataFrame([nuevo_registro])
                df_google = pd.concat([df_google, nuevo_df], ignore_index=True)
                logging.info(f"Agregado nuevo registro para {regional} - {fecha_reporte_str}")
        else:
            # Primera vez o DataFrame vacío, agregar el registro
            df_google = pd.DataFrame([nuevo_registro])
        
        # 6. Escribir datos de vuelta a Google Sheets
        valores_actualizados = [df_google.columns.tolist()] + df_google.where(pd.notna(df_google), None).values.tolist()
        
        # Redimensionar la hoja antes de escribir para liberar celdas del limite de 10M
        redimensionar_hoja_google_sheets(client, sheet_id, sheet_name, len(valores_actualizados), len(valores_actualizados[0]))
        
        client.spreadsheets().values().clear(spreadsheetId=sheet_id, range=sheet_name).execute()
        body = {'values': valores_actualizados}
        client.spreadsheets().values().update(
            spreadsheetId=sheet_id,
            range=f"{sheet_name}!A1",
            valueInputOption='USER_ENTERED',
            body=body
        ).execute()
        
        logging.info(f"Porcentajes de sincronización subidos correctamente a '{sheet_name}'")
        return True
        
    except Exception as e:
        logging.error(f"Error al subir porcentajes de sincronización: {e}")
        return False


def obtener_promedios_mensuales_sheets(regional):
    """
    Consulta la hoja 'PorcSinc' en Google Sheets y calcula el promedio
    de los porcentajes para el mes y año actual.
    
    Returns:
        tuple: (promedio_al_dia, promedio_revision, num_ejecuciones)
    """
    try:
        # 1. Conexión a Google Sheets
        creds_path = resource_path('credentials_gebilo.json')
        if not os.path.exists(creds_path):
            return None, None, 0
        
        scope = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_file(creds_path, scopes=scope)
        client = build('sheets', 'v4', credentials=creds)
        sheet_id = '12GPBc0JAMRY2b7YMWcFDGiAP9xOrfDNgoZdiwQyuWdc'
        sheet_name = 'PorcSinc'
        
        # 2. Leer datos existentes
        range_name = f"{sheet_name}!A:G"
        result = client.spreadsheets().values().get(spreadsheetId=sheet_id, range=range_name).execute()
        values = result.get('values', [])
        
        if not values or len(values) <= 1:
            return None, None, 0
            
        header = values[0]
        df_google = pd.DataFrame(values[1:], columns=header)
        
        # 3. Filtrar por Regional y Mes/Año Actual
        ahora = datetime.now()
        mes_actual = ahora.month
        año_actual = ahora.year
        
        df_google['FechaDT'] = pd.to_datetime(df_google['Fecha'], errors='coerce')
        
        mask = (
            (df_google['Regional'].astype(str) == str(regional)) &
            (df_google['FechaDT'].dt.month == mes_actual) &
            (df_google['FechaDT'].dt.year == año_actual)
        )
        
        df_mes = df_google[mask].copy()
        
        if df_mes.empty:
            return None, None, 0
            
        # 4. Calcular promedios
        # Convertir a numérico por si vienen como strings desde Sheets
        df_mes['Porc. al día'] = pd.to_numeric(df_mes['Porc. al día'], errors='coerce')
        df_mes['Porc. Revisión'] = pd.to_numeric(df_mes['Porc. Revisión'], errors='coerce')
        
        promedio_al_dia = round(df_mes['Porc. al día'].mean(), 1)
        promedio_revision = round(df_mes['Porc. Revisión'].mean(), 1)
        num_ejecuciones = len(df_mes)
        
        return promedio_al_dia, promedio_revision, num_ejecuciones
        
    except Exception as e:
        logging.error(f"Error al obtener promedios desde Sheets: {e}")
        return None, None, 0


def enviar_correo_notificacion(asunto, cuerpo_html, destinatarios, archivo_adjunto=None):
    """
    Se conecta al servidor SMTP y envía un correo HTML, opcionalmente con un archivo o archivos adjuntos.

    Args:
        asunto (str): El asunto del correo.
        cuerpo_html (str): El contenido del correo en formato HTML.
        destinatarios (list): Una lista de las direcciones de correo de los destinatarios.
        archivo_adjunto (str/list/tuple, optional): La ruta o lista de rutas completas a los archivos a adjuntar. Por defecto es None.
    """
    # --- Credenciales configuradas ---
    remitente = "victorpty999@gmail.com"
    password_app = "tcdlqiyfgoirkruw"
    servidor_smtp = "smtp.gmail.com"
    puerto = 587
    
    if not destinatarios:
        logging.warning("No se enviaron correos: no hay destinatarios configurados.")
        return False

    # --- Construir el mensaje ---
    msg = EmailMessage()
    msg['Subject'] = asunto
    msg['From'] = remitente
    msg['To'] = ", ".join(destinatarios)
    #msg['Cc'] = "gebilo@infoplazas.org.pa" # <-- Línea añadida para la copia
    msg.set_content("Este correo contiene formato HTML. Por favor, actívelo para ver el contenido.")
    msg.add_alternative(cuerpo_html, subtype='html')
    
    # --- Lógica para adjuntar archivo(s) ---
    if archivo_adjunto:
        archivos = [archivo_adjunto] if isinstance(archivo_adjunto, str) else archivo_adjunto
        for adjunto in archivos:
            if adjunto and os.path.exists(adjunto):
                try:
                    with open(adjunto, 'rb') as f:
                        archivo_data = f.read()
                        archivo_nombre = os.path.basename(adjunto)
                    # Adjunta el archivo al mensaje
                    msg.add_attachment(archivo_data, maintype='application', subtype='octet-stream', filename=archivo_nombre)
                    logging.info(f"Archivo adjuntado al correo: {archivo_nombre}")
                except Exception as e:
                    logging.error(f"No se pudo adjuntar el archivo {adjunto}: {e}")

    MAX_RETRIES = 10
    RETRY_DELAY = 60

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            # --- Conectar y enviar ---
            with smtplib.SMTP(servidor_smtp, puerto) as smtp:
                smtp.starttls()
                smtp.login(remitente, password_app)
                smtp.send_message(msg)
            logging.info(f"Correo de notificación enviado a: {', '.join(destinatarios)}")
            return True

        except smtplib.SMTPAuthenticationError as e:
            logging.error(f"Error FATAL de autenticación SMTP: {e}. Verifique usuario/crédenciales. No se reintentará.")
            break

        except (socket.gaierror, ConnectionRefusedError, smtplib.SMTPConnectError, smtplib.SMTPServerDisconnected, OSError) as e:
            logging.warning(f"Error de conexión enviando correo (Intento {attempt}/{MAX_RETRIES}): {e}")
            if attempt < MAX_RETRIES:
                logging.info(f"Esperando {RETRY_DELAY} segundos antes del reintento...")
                time.sleep(RETRY_DELAY)
            else:
                logging.error(f"Se agotaron los {MAX_RETRIES} intentos de envío. No se pudo enviar el correo.")
        
        except Exception as e:
            logging.error(f"Error inesperado no manejado al enviar correo: {e}")
            break

    return False



def ejecutar_tarea_automatica(app_instance, regiones_seleccionadas):
    logging.info(f"PILOTO AUTOMÁTICO: Iniciando ciclo para las regionales: {', '.join(regiones_seleccionadas)}")
    
    # Obtiene una ruta de ejemplo desde la UI para determinar la carpeta principal
    sample_path = app_instance.carpeta_entrada.get()
    if not sample_path:
        logging.error("No se ha seleccionado una carpeta principal para las regionales.")
        # Opcional: Podrías notificar al usuario aquí, pero como es automático, el log es lo principal.
        return

    # --- CORRECTION HERE ---
    # We get the parent directory of the selected path.
    master_path = os.path.dirname(sample_path)

     # Bucle principal para procesar cada regional seleccionada
    # Bucle principal para procesar cada regional seleccionada
    for nombre_regional in regiones_seleccionadas:
        if not app_instance.piloto_activo:
            logging.info("El piloto automático fue detenido a mitad de ciclo.")
            break 

        logging.info(f"Procesando regional: {nombre_regional}")

        folder_name = REGION_FOLDER_MAP.get(nombre_regional)
        if not folder_name:
            logging.error(f"No se encontró el nombre de la carpeta para la regional: {nombre_regional}")
            continue

        ruta_carpeta_regional = os.path.join(master_path, folder_name)
        app_instance.carpeta_entrada.set(ruta_carpeta_regional)

        nombre_archivo = f"Reporte_{nombre_regional.replace(' ', '_')}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

        # Creamos el hilo de procesamiento
        hilo_proceso = threading.Thread(target=app_instance._proceso_thread, 
                                        args=('Automático', True, nombre_archivo), 
                                        daemon=True)
        # Lo iniciamos
        hilo_proceso.start()
        # Con .join(), el bucle principal se detiene aquí hasta que el hilo termine
        hilo_proceso.join()

        logging.info(f"Procesamiento de {nombre_regional} completado. Pausa de 10 segundos.")
        time.sleep(10)

    # Fin del bucle principal



def calcular_proxima_ejecucion(dias_seleccionados_str):
    """
    Calcula la próxima ejecución del piloto automático.
    
    Lógica:
    1. Calcula la próxima ejecución según días seleccionados
    2. Verifica si el último día del mes está antes de esa fecha
    3. Si es así, programa para el último día del mes
    4. Retorna (fecha_proxima, es_fin_de_mes)
    """
    ahora = datetime.now()
    mapa_dias = {"Lunes": 0, "Martes": 1, "Miércoles": 2, "Jueves": 3, "Viernes": 4, "Sábado": 5, "Domingo": 6}
    dias_seleccionados_num = sorted([mapa_dias[d] for d in dias_seleccionados_str])
    
    # Calcular próxima ejecución según días seleccionados
    if not dias_seleccionados_num:
        proxima_normal = (ahora + timedelta(days=1)).replace(hour=8, minute=0, second=0, microsecond=0)
    else:
        dia_actual_num = ahora.weekday()
        proxima_normal = None
        
        for dia_semana in dias_seleccionados_num:
            if dia_semana > dia_actual_num or (dia_semana == dia_actual_num and ahora.hour < 8):
                dias_a_sumar = dia_semana - dia_actual_num
                proxima_normal = (ahora + timedelta(days=dias_a_sumar)).replace(hour=8, minute=0, second=0, microsecond=0)
                break
        
        if proxima_normal is None:
            dias_a_sumar = (7 - dia_actual_num) + dias_seleccionados_num[0]
            proxima_normal = (ahora + timedelta(days=dias_a_sumar)).replace(hour=8, minute=0, second=0, microsecond=0)
    
    # Calcular el último día del mes actual
    # Obtener el primer día del próximo mes
    if ahora.month == 12:
        primer_dia_proximo_mes = datetime(ahora.year + 1, 1, 1)
    else:
        primer_dia_proximo_mes = datetime(ahora.year, ahora.month + 1, 1)
    
    # Restar un día para obtener el último día del mes actual
    ultimo_dia_mes = primer_dia_proximo_mes - timedelta(days=1)
    ultimo_dia_mes_ejecutable = ultimo_dia_mes.replace(hour=8, minute=0, second=0, microsecond=0)
    
    # Verificar si ya pasamos el último día del mes o si ya es hoy y ya pasaron las 8 AM
    if ahora.date() > ultimo_dia_mes.date() or (ahora.date() == ultimo_dia_mes.date() and ahora.hour >= 8):
        # Ya pasó el último día del mes, retornar la próxima ejecución normal
        return (proxima_normal, False)
    
    # Comparar: ¿El último día del mes viene antes que la próxima ejecución normal?
    if ultimo_dia_mes_ejecutable < proxima_normal:
        # Ejecutar en el último día del mes
        return (ultimo_dia_mes_ejecutable, True)
    elif ultimo_dia_mes_ejecutable == proxima_normal:
        # Coinciden, ejecutar solo una vez pero marcar como fin de mes
        return (proxima_normal, True)
    else:
        # La próxima ejecución normal viene primero
        return (proxima_normal, False)

def bucle_piloto_automatico(app_instance, dias_semana, regiones_seleccionadas, proxima_ejecucion_guardada=None, run_immediately=True):
    if run_immediately:
        # Pasa la lista de regionales
        ejecutar_tarea_automatica(app_instance, regiones_seleccionadas)

    proxima = proxima_ejecucion_guardada
    while app_instance.piloto_activo:
        if not proxima:
            # Obtener próxima ejecución (ahora retorna tupla)
            proxima, es_fin_mes = calcular_proxima_ejecucion(dias_semana)
        else:
            # Si venimos de una ejecución guardada, asumir que no es fin de mes
            es_fin_mes = False
            
        app_instance.proxima_ejecucion_dt = proxima
        
        # Mensaje de log diferenciado
        if es_fin_mes:
            logging.info(f"PILOTO AUTOMÁTICO: Próxima ejecución: {proxima} (Último día del mes)")
        else:
            logging.info(f"PILOTO AUTOMÁTICO: Próxima ejecución: {proxima}")
            
        app_instance.after(0, lambda p=proxima: app_instance.actualizar_cronometro(p))
        
        while datetime.now() < proxima and app_instance.piloto_activo:
            time.sleep(5)
            
        if app_instance.piloto_activo:
            # Pasa la lista de regionales
            ejecutar_tarea_automatica(app_instance, regiones_seleccionadas)
            proxima = None  # Resetear para calcular la siguiente
    
    logging.info("PILOTO AUTOMÁTICO: Bucle finalizado.")


class AnimationWindow(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent; self.width, self.height = 400, 200
        parent_x, parent_y, parent_w, parent_h = parent.winfo_x(), parent.winfo_y(), parent.winfo_width(), parent.winfo_height()
        x, y = parent_x + (parent_w // 2) - (self.width // 2), parent_y + (parent_h // 2) - (self.height // 2)
        self.geometry(f"{self.width}x{self.height}+{x}+{y}")
        self.overrideredirect(True); self.transient(parent); self.grab_set()
        try: self.attributes("-alpha", 0.95)
        except tk.TclError: pass
        self.config(bg="#2B2B2B"); self.canvas = tk.Canvas(self, width=self.width, height=130, bg="#2B2B2B", highlightthickness=0); self.canvas.pack(pady=(20, 0))
        self.status_var = tk.StringVar(value="Iniciando..."); ttk.Label(self, textvariable=self.status_var, foreground="#FFFFFF", background="#2B2B2B", font=("Helvetica", 10)).pack(pady=5)
        self.angle, self.running = 0, False; self.circles = []
        colors = ["#08D9D6", "#252A34", "#FF2E63", "#EAEAEA", "#00ADB5"]
        for i in range(10): self.circles.append([10 + i * 8, i * (2 * math.pi / 10), colors[i % len(colors)]])
    def start_animation(self): self.running = True; self.animate()
    def stop_animation(self): self.running = False; self.grab_release(); self.destroy()
    def animate(self):
        if not self.running: return
        self.canvas.delete("all"); center_x, center_y = self.width / 2, self.height / 2 - 20
        for i, (radius, angle_offset, color) in enumerate(self.circles):
            y_offset = 30 * math.sin(self.angle/2 + i); x = center_x + (radius + y_offset) * math.cos(self.angle + angle_offset); y = center_y + (radius + y_offset) * math.sin(self.angle + angle_offset)
            r = 7; self.canvas.create_oval(x - r, y - r, x + r, y + r, fill=color, outline="")
        self.angle += 0.05; self.after(25, self.animate)
    def update_status(self, text): self.status_var.set(text); self.update_idletasks()

class TreeviewToolTip:
    """Implementa un Tooltip para filas de un Treeview."""
    def __init__(self, tree, text):
        self.tree = tree
        self.text = text
        self.tipwindow = None
        self.last_item = None
        self.tree.bind("<Motion>", self.check_tooltip)
        self.tree.bind("<Leave>", self.hide_tooltip)

    def check_tooltip(self, event):
        item = self.tree.identify_row(event.y)
        if item:
            if item != self.last_item:
                self.last_item = item
                self.show_tooltip(event.x_root + 15, event.y_root + 10)
        else:
             self.hide_tooltip(None)

    def show_tooltip(self, x, y):
         if self.tipwindow: self.tipwindow.destroy()
         self.tipwindow = tk.Toplevel(self.tree)
         self.tipwindow.wm_overrideredirect(True)
         self.tipwindow.wm_geometry(f"+{x}+{y}")
         label = tk.Label(self.tipwindow, text=self.text, justify=tk.LEFT,
                          background="#ffffe0", relief=tk.SOLID, borderwidth=1,
                          font=("tahoma", "8", "normal"))
         label.pack(ipadx=1)

    def hide_tooltip(self, event):
        if self.tipwindow:
            self.tipwindow.destroy()
            self.tipwindow = None
        self.last_item = None


class App(ttk.Window):

    def construir_cuerpo_correo_de_fallo(self, regional, periodo, origen, error_msg):
        """Construye un correo simple para notificar un error en el proceso."""
        
        # Saludo
        saludo = "Buen día" if datetime.now().hour < 12 else "Buenas tardes"
        
        # Cuerpo del correo de error
        cuerpo_completo = f"""
        <html><head><style>
            body {{ font-family: sans-serif; font-size: 14px; }}
            strong {{ color: #d9534f; }}
        </style></head><body>
            <p>{saludo},</p>
            <p>Se ha producido un error durante la ejecución del analizador de bases de datos.</p>
            <h3>Detalles del Fallo:</h3>
            <ul>
                <li><strong>Origen de Ejecución:</strong> {origen}</li>
                <li><strong>Regional Afectada:</strong> {regional}</li>
                <li><strong>Periodo Analizado:</strong> {periodo}</li>
                <li><strong>Mensaje de Error:</strong> {error_msg}</li>
            </ul>
            <p>Por favor, revise el archivo <code>app.log</code> para más detalles técnicos.</p>
        </body></html>
        """
        return cuerpo_completo

    def __init__(self):
        super().__init__(themename="darkly", title=f"Analizador de Bases de Datos by VicTor - v{VERSION}")
        self.state('zoomed')
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        try:
            self.iconbitmap(resource_path('VicTor.ico'))
            
        except:
            logging.warning("No se pudo cargar el ícono 'VicTor.ico'")
            
        
        # Variables de estado
        self.carpeta_entrada = tk.StringVar()
        self.resumen = pd.DataFrame()
        self.resumen_cruzado = pd.DataFrame()
        self.piloto_activo = False
        self.hilo_piloto = None
        self.animation_window = None
        self.proxima_ejecucion_var = tk.StringVar()
        self.proxima_ejecucion_dt = None
        self.execution_history = []
        self.last_manual_run_id = None
        self.recipients_list = [] 
        self.editing_email = None # Variable para rastrear qué correo se está editando 
        self.supabase_url = "https://jcozaaifpfukqlypfuqq.supabase.co"
        self.supabase_key = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6Impjb3phYWlmcGZ1a3FseXBmdXFxIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4MDY4Mjg4MiwiZXhwIjoyMDk2MjU4ODgyfQ.80jkincEydxGbuhqKYEjJBOf1LJ3TkhthVgl0sPah8k"

        self.create_widgets()

    def create_widgets(self):
        main_frame = ttk.Frame(self, padding=15)
        main_frame.pack(expand=True, fill=BOTH)
        main_frame.grid_rowconfigure(1, weight=1)
        main_frame.grid_columnconfigure(0, weight=1)

        top_container = ttk.Frame(main_frame)
        top_container.pack(fill=X, pady=(0, 15))
        top_container.grid_columnconfigure(0, weight=1)
        top_container.grid_columnconfigure(1, weight=1)

        # --- Panel de Ejecución Manual ---
        manual_frame = ttk.LabelFrame(top_container, text=" Ejecución Manual y Configuración Principal ", padding=15)
        manual_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        manual_frame.grid_columnconfigure(1, weight=1)
        
        # La etiqueta ahora indica que es la carpeta principal
        ttk.Label(manual_frame, text="Carpeta Principal de Regionales:").grid(row=0, column=0, padx=5, pady=5, sticky=W)
        ttk.Entry(manual_frame, textvariable=self.carpeta_entrada, width=40).grid(row=0, column=1, padx=5, pady=5, sticky=EW)
        ttk.Button(manual_frame, text="Examinar", command=self.seleccionar_carpeta, bootstyle="info-outline-toolbutton").grid(row=0, column=2, padx=5, pady=5)

        ttk.Label(manual_frame, text="Fecha Inicial:").grid(row=1, column=0, padx=5, pady=10, sticky=W)
        self.date_start = DateEntry(manual_frame, bootstyle=PRIMARY, dateformat='%d/%m/%Y')
        self.date_start.grid(row=1, column=1, padx=5, pady=5, sticky=W)
        ttk.Label(manual_frame, text="Fecha Final:").grid(row=2, column=0, padx=5, pady=5, sticky=W)
        self.date_end = DateEntry(manual_frame, bootstyle=PRIMARY, dateformat='%d/%m/%Y')
        self.date_end.grid(row=2, column=1, padx=5, pady=5, sticky=W)
        
        btn_frame = ttk.Frame(manual_frame)
        btn_frame.grid(row=3, column=0, columnspan=3, pady=10)
        for text, period in [("Este mes", "Este mes"), ("Mes anterior", "Mes anterior"), ("Este año", "Este año"), ("Año anterior", "Año anterior")]:
            ttk.Button(btn_frame, text=text, command=lambda p=period: self.establecer_fechas(p), bootstyle="info-outline-toolbutton").pack(side=LEFT, padx=3)
        
        action_frame = ttk.Frame(manual_frame)
        action_frame.grid(row=4, column=0, columnspan=3, pady=15)
        ttk.Button(action_frame, text="📊 Generar Reporte", command=lambda: self.ejecutar_proceso(origen='Manual'), bootstyle=PRIMARY).pack(pady=5)
        self.btn_subir_g = ttk.Button(action_frame, text="☁️ Subir Resumen a G. Sheets", command=self.subir_resumen_a_sheets, bootstyle=INFO, state=DISABLED)
        self.btn_subir_g.pack(pady=5)
        
        # Botón para Analítica Histórica
        ttk.Button(action_frame, text="📈 Ver Historial Sucursales", command=lambda: sucursales_history.mostrar_ventana_historial(self), bootstyle="secondary-outline").pack(pady=5)


        # --- Panel de Piloto Automático ---
        piloto_frame = ttk.LabelFrame(top_container, text=" Piloto Automático ", padding=15)
        piloto_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
        self.piloto_switch = ttk.Checkbutton(piloto_frame, text="Activar Piloto Automático", bootstyle="success-square-toggle", command=self.toggle_piloto_frame)
        self.piloto_switch.pack(anchor=W)
        self.piloto_options_frame = ttk.Frame(piloto_frame)
        self.piloto_options_frame.pack(fill=X, pady=10, padx=5)
        ttk.Label(self.piloto_options_frame, text="Ejecutar los días:").pack(side=LEFT, padx=5)
        self.frame_dias_semana = ttk.Frame(self.piloto_options_frame)
        self.frame_dias_semana.pack(side=LEFT, padx=10)
        self.vars_dias = {dia: tk.BooleanVar() for dia in ["L", "M", "X", "J", "V", "S", "D"]}
        for dia, var in self.vars_dias.items():
            ttk.Checkbutton(self.frame_dias_semana, text=dia, variable=var, bootstyle="info-outline-toolbutton").pack(side=LEFT)

        # --- Selección de Regionales ---
        regionales_frame = ttk.LabelFrame(piloto_frame, text=" Regionales a Procesar ", padding=10)
        regionales_frame.pack(fill=X, expand=True, pady=(10,0))
        self.regional_vars = {}
        for regional_name in REGION_FOLDER_MAP.keys():
            var = tk.BooleanVar()
            self.regional_vars[regional_name] = var
            ttk.Checkbutton(regionales_frame, text=regional_name, variable=var, bootstyle="primary-round-toggle").pack(side=LEFT, padx=10, pady=5)

        self.btn_iniciar_piloto = ttk.Button(piloto_frame, text="Iniciar Piloto", bootstyle="success", command=self.iniciar_piloto)
        self.btn_iniciar_piloto.pack(pady=(15,0))
        self.piloto_options_frame.pack_forget()

        # --- Pestañas ---
        self.notebook = ttk.Notebook(main_frame, bootstyle="dark")
        self.notebook.pack(fill=BOTH, expand=True, padx=10, pady=10)

        historial_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(historial_tab, text=' Historial de Ejecuciones ')
        self.create_history_table(historial_tab)

        self.incidentes_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.incidentes_tab, text=' Sincronización ')
        self.create_incidents_table(self.incidentes_tab)

        destinatarios_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(destinatarios_tab, text=' 📧 Destinatarios ')
        self.create_recipients_tab(destinatarios_tab)

        # --- Pie de página ---
        footer = ttk.Frame(self, padding=5)
        footer.pack(side=BOTTOM, fill=X)
        self.status_label = ttk.Label(footer, text="Listo.", anchor=W)
        self.status_label.pack(side=LEFT, padx=10)
        self.cronometro_label = ttk.Label(footer, textvariable=self.proxima_ejecucion_var)
        self.cronometro_label.pack(side=LEFT, padx=10)
        ttk.Label(footer, text=AUTOR, anchor=E).pack(side=RIGHT, padx=10)

    def create_history_table(self, parent):
        cols = ('id', 'fecha_hora', 'origen', 'regional', 'periodo', 'archivos', 'registros', 'proceso', 'subida', 'correo_enviado', 'incidencias')
        self.history_tree = ttk.Treeview(parent, columns=cols, show='headings', bootstyle=DARK)
        headings = {'id': '#', 'fecha_hora': 'Fecha y Hora', 'origen': 'Origen', 'regional': 'Regional', 'periodo': 'Periodo', 'archivos': 'Archivos', 'registros': 'Registros', 'proceso': 'Completo', 'subida': 'Sheet/Supa', 'correo_enviado': 'Correo Enviado', 'incidencias': 'Incidencias'}
        for col in cols: self.history_tree.heading(col, text=headings[col], command=lambda c=col: self.sort_treeview(self.history_tree, c, False))
        
        widths = {'id': 40, 'fecha_hora': 140, 'origen': 80, 'regional': 100, 'periodo': 160, 'archivos': 60, 'registros': 70, 'proceso': 70, 'subida': 80, 'correo_enviado': 80, 'incidencias': 200}
        for col, width in widths.items(): self.history_tree.column(col, width=width, anchor=CENTER if col not in ['incidencias'] else W)
        
        self.history_tree.pack(side=LEFT, fill=BOTH, expand=True)
        scrollbar = ttk.Scrollbar(parent, orient=VERTICAL, command=self.history_tree.yview); self.history_tree.configure(yscroll=scrollbar.set); scrollbar.pack(side=RIGHT, fill=Y)
        self.history_tree.tag_configure('oddrow', background='#303030'); self.history_tree.tag_configure('evenrow', background='#2B2B2B')
   

    def create_incidents_table(self, parent):
        cols = ('fecha_sinc', 'carpeta', 'fecha_archivo', 'dias_sin_sinc', 'primer_reg', 'ultimo_reg', 'observacion')
        self.incidents_tree = ttk.Treeview(parent, columns=cols, show='headings', bootstyle=DARK)
        
        headings = {
            'fecha_sinc': 'Fecha Sincronización', 
            'carpeta': 'Sucursal', 
            'fecha_archivo': 'Fecha Archivo', 
            'dias_sin_sinc': 'Días sin Sinc.',
            'primer_reg': 'Primer Registro', 
            'ultimo_reg': 'Último Registro', 
            'observacion': 'Observación'
        }
        for col in cols:
            self.incidents_tree.heading(col, text=headings[col], command=lambda c=col: self.sort_treeview(self.incidents_tree, c, False))
        
        self.incidents_tree.column('fecha_sinc', width=140, anchor=W)
        self.incidents_tree.column('carpeta', width=120, anchor=W)
        self.incidents_tree.column('fecha_archivo', width=100, anchor=CENTER)
        self.incidents_tree.column('dias_sin_sinc', width=100, anchor=CENTER)
        
        self.incidents_tree.pack(side=LEFT, fill=BOTH, expand=True)
        scrollbar = ttk.Scrollbar(parent, orient=VERTICAL, command=self.incidents_tree.yview)
        self.incidents_tree.configure(yscroll=scrollbar.set)
        scrollbar.pack(side=RIGHT, fill=Y)
        self.incidents_tree.tag_configure('oddrow', background='#303030')
        self.incidents_tree.tag_configure('evenrow', background='#2B2B2B')

    def create_recipients_tab(self, parent):
        """Crea el contenido de la pestaña de gestión de destinatarios con una fila horizontal."""
        
        # Frame contenedor
        container = ttk.Frame(parent)
        container.pack(fill=BOTH, expand=True)

        # Frame para el formulario (Fila Única)
        form_frame = ttk.LabelFrame(container, text=" Gestionar Destinatario ", padding=10)
        form_frame.pack(fill=X, pady=(0, 10))

        # --- Fila de Controles ---
        # Nombre
        ttk.Label(form_frame, text="Nombre:").pack(side=LEFT, padx=(0, 5))
        self.recipient_name = ttk.Entry(form_frame, width=25)
        self.recipient_name.pack(side=LEFT, padx=(0, 10), fill=X, expand=True)

        # Correo
        ttk.Label(form_frame, text="Correo:").pack(side=LEFT, padx=(0, 5))
        self.recipient_email = ttk.Entry(form_frame, width=30)
        self.recipient_email.pack(side=LEFT, padx=(0, 10), fill=X, expand=True)

        # Cargo (Opciones Extendidas)
        ttk.Label(form_frame, text="Cargo:").pack(side=LEFT, padx=(0, 5))
        cargos = ["Facilitador", "Enlace", "Supervisor", "Administrador", "Director", "Subdirector", "Asistente"]
        self.recipient_cargo = ttk.Combobox(form_frame, values=cargos, state="readonly", width=20)
        self.recipient_cargo.pack(side=LEFT, padx=(0, 10))

        # Regional (Opción "Todas" agregada)
        ttk.Label(form_frame, text="Regional:").pack(side=LEFT, padx=(0, 5))
        regionales = ["Todas", "Los Santos", "Chiriquí", "Veraguas", "Panamá"]
        self.recipient_regional = ttk.Combobox(form_frame, values=regionales, state="readonly", width=20)
        self.recipient_regional.pack(side=LEFT, padx=(0, 10))

        # Botones de Acción (En la misma línea)
        self.btn_add_recipient = ttk.Button(form_frame, text="Añadir", bootstyle="success", command=self.add_recipient, width=15)
        self.btn_add_recipient.pack(side=LEFT, padx=(5, 5))
        
        self.btn_cancel_edit = ttk.Button(form_frame, text="Cancelar", bootstyle="success", command=self.cancel_edit, width=15)
        self.btn_cancel_edit.pack(side=LEFT)
        self.btn_cancel_edit.pack_forget() # Oculto por defecto

        ttk.Button(form_frame, text="Quitar", bootstyle="danger-outline", command=self.remove_recipient, width=15).pack(side=LEFT, padx=(5, 0))

        # --- Tabla de Destinatarios ---
        cols = ('nombre', 'correo', 'cargo', 'regional', 'activo')
        self.recipients_tree = ttk.Treeview(container, columns=cols, show='headings', bootstyle=DARK)
        
        headings = {'nombre': 'Nombre', 'correo': 'Correo', 'cargo': 'Cargo', 'regional': 'Regional', 'activo': 'Activo'}
        for col in cols: self.recipients_tree.heading(col, text=headings[col], command=lambda c=col: self.sort_treeview(self.recipients_tree, c, False))
        
        self.recipients_tree.column('nombre', width=150)
        self.recipients_tree.column('correo', width=150)
        self.recipients_tree.column('cargo', width=130)
        self.recipients_tree.column('regional', width=130)
        self.recipients_tree.column('activo', width=80, anchor='center')

        self.recipients_tree.pack(side=LEFT, fill=BOTH, expand=True)
        scrollbar = ttk.Scrollbar(container, orient=VERTICAL, command=self.recipients_tree.yview)
        self.recipients_tree.configure(yscroll=scrollbar.set)
        scrollbar.pack(side=RIGHT, fill=Y)

        # Evento Doble Clic para Editar
        self.recipients_tree.bind("<Double-1>", self.load_recipient_for_edit)
        self.recipients_tree.bind("<Delete>", lambda event: self.remove_recipient())
        self.recipients_tree.bind("<ButtonRelease-1>", self.on_recipients_click)

        
        # --- Tooltip de ayuda ---
        TreeviewToolTip(self.recipients_tree, "Doble Clic: Editar\nSuprimir: Eliminar")

    def load_recipient_for_edit(self, event):
        """Carga los datos del destinatario seleccionado en el formulario para editar."""
        column = self.recipients_tree.identify_column(event.x)
        if column == "#5":
            return
        selected_item = self.recipients_tree.focus()
        if not selected_item: return

        values = self.recipients_tree.item(selected_item, 'values')
        if not values: return

        # Cargar valores en inputs
        self.recipient_name.delete(0, END); self.recipient_name.insert(0, values[0])
        self.recipient_email.delete(0, END); self.recipient_email.insert(0, values[1])
        self.recipient_cargo.set(values[2])
        self.recipient_regional.set(values[3])

        # Estado de edición
        self.editing_email = values[1] # Usamos el correo original como ID
        self.btn_add_recipient.config(text="Guardar Cambios", bootstyle="warning")
        self.btn_cancel_edit.pack(side=LEFT, before=self.btn_add_recipient) # Mostrar botón cancelar a la izq del guardar? O next.
        # Mejor pack de nuevo para orden visual si es necesario, pero simple pack/unpack funciona.
        self.btn_cancel_edit.pack(side=LEFT, after=self.btn_add_recipient, padx=5)

    def on_recipients_click(self, event):
        region = self.recipients_tree.identify("region", event.x, event.y)
        if region == "cell":
            column = self.recipients_tree.identify_column(event.x)
            if column == "#5":
                item_id = self.recipients_tree.identify_row(event.y)
                if item_id:
                    values = list(self.recipients_tree.item(item_id, 'values'))
                    if len(values) >= 5:
                        current_activo = values[4]
                        new_activo = "☐" if current_activo == "☑" else "☑"
                        values[4] = new_activo
                        self.recipients_tree.item(item_id, values=values)
                        
                        # Actualizar la lista en memoria
                        email = values[1]
                        for r in self.recipients_list:
                            if r['correo'] == email:
                                r['activo'] = (new_activo == "☑")
                                break
                        self.save_settings()

    def cancel_edit(self):
        """Cancela el modo edición y limpia el formulario."""
        self.recipient_name.delete(0, END)
        self.recipient_email.delete(0, END)
        self.recipient_cargo.set('')
        self.recipient_regional.set('')
        self.editing_email = None
        self.btn_add_recipient.config(text="Añadir", bootstyle="success")
        self.btn_cancel_edit.pack_forget()

    def add_recipient(self):
        """Añade o Actualiza un destinatario."""
        name = self.recipient_name.get().strip()
        email = self.recipient_email.get().strip()
        cargo = self.recipient_cargo.get()
        regional = self.recipient_regional.get()

        if not (name and email and cargo and regional):
            Messagebox.show_warning("Todos los campos son requeridos.", "Faltan Datos", parent=self)
            return

        # --- MODO EDICIÓN ---
        if self.editing_email:
            # 1. Actualizar en la lista interna
            found = False
            for r in self.recipients_list:
                if r['correo'] == self.editing_email:
                    r['nombre'] = name
                    r['correo'] = email
                    r['cargo'] = cargo
                    r['regional'] = regional
                    found = True
                    break
            
            # Si cambió el correo, verificar duplicados (opcional, pero buena práctica)
            # En este caso simple, asumimos que si edita su propio correo está bien.

            # 2. Actualizar en Treeview
            # Buscar el item que tiene el editing_email viejo
            for item in self.recipients_tree.get_children():
                vals = self.recipients_tree.item(item, 'values')
                if vals[1] == self.editing_email:
                    activo_str = vals[4] if len(vals) >= 5 else "☑"
                    self.recipients_tree.item(item, values=(name, email, cargo, regional, activo_str))
                    break
            
            # Salir de modo edición
            self.cancel_edit()

        # --- MODO AÑADIR ---
        else:
            # Verificar duplicados por correo
            if any(r['correo'] == email for r in self.recipients_list):
                 Messagebox.show_error("Este correo ya está registrado.", "Duplicado", parent=self)
                 return

            new_recipient = {"nombre": name, "correo": email, "cargo": cargo, "regional": regional, "activo": True}
            self.recipients_list.append(new_recipient)
            self.recipients_tree.insert('', END, values=(name, email, cargo, regional, "☑"))
            
            # Limpiar
            self.recipient_name.delete(0, END)
            self.recipient_email.delete(0, END)
            self.recipient_cargo.set('')
            self.recipient_regional.set('')

        # Guardar cambios sin notificación de éxito
        self.save_settings()

    def remove_recipient(self):
        """Quita el destinatario seleccionado con confirmación."""
        selected_item = self.recipients_tree.focus()
        if not selected_item: return
            
        # Confirmación
        confirm = Messagebox.show_question("¿Está seguro de eliminar este destinatario?", "Confirmar Eliminación", parent=self)
        if confirm != 'Yes': return # 'Yes' es el retorno estándar de show_question en ttkbootstrap/tkinter

        selected_values = self.recipients_tree.item(selected_item, 'values')
        email_to_remove = selected_values[1] 
        
        self.recipients_list = [r for r in self.recipients_list if r['correo'] != email_to_remove]
        self.recipients_tree.delete(selected_item)
        
        # Si estaba editando justo el que borró (caso raro), cancelar edit
        if self.editing_email == email_to_remove:
            self.cancel_edit()

        self.save_settings()

    def sort_treeview(self, tree, col, reverse):
        data = [(tree.set(item, col), item) for item in tree.get_children('')]
        try: data.sort(key=lambda t: float(t[0]), reverse=reverse)
        except (ValueError, TypeError): data.sort(key=lambda t: str(t[0]), reverse=reverse)
        for index, (val, item) in enumerate(data): tree.move(item, '', index)
        tree.heading(col, command=lambda: self.sort_treeview(tree, col, not reverse))
    
    def on_closing(self): self.save_settings(); self.destroy()

    def seleccionar_carpeta(self):
        path = filedialog.askdirectory(title="Seleccione la Carpeta Raíz");
        if path: self.carpeta_entrada.set(path); self.save_settings()

    def establecer_fechas(self, tipo):
        hoy = datetime.today()
        if tipo == "Este mes": start, end = hoy.replace(day=1), hoy
        elif tipo == "Mes anterior": end = hoy.replace(day=1) - timedelta(days=1); start = end.replace(day=1)
        elif tipo == "Este año": start, end = hoy.replace(month=1, day=1), hoy
        else: end = hoy.replace(year=hoy.year-1, month=12, day=31); start = end.replace(month=1, day=1)
        self.date_start.entry.delete(0, END); self.date_start.entry.insert(0, start.strftime('%d/%m/%Y'))
        self.date_end.entry.delete(0, END); self.date_end.entry.insert(0, end.strftime('%d/%m/%Y'))
        self.save_settings()

    def toggle_piloto_frame(self):
        if self.piloto_switch.instate(['selected']): self.piloto_options_frame.pack(fill=X, pady=10, padx=5)
        else: self.piloto_options_frame.pack_forget(); self.detener_piloto() if self.piloto_activo else None

    def show_animation(self, status_text="Procesando..."):
        if self.animation_window and self.animation_window.winfo_exists():
            self.animation_window.update_status(status_text)
            return
        self.bind("<FocusIn>", self._lift_animation_window)
        self.animation_window = AnimationWindow(self)
        self.animation_window.update_status(status_text)
        self.after(0, self.animation_window.start_animation)

    def hide_animation(self):
        if self.animation_window:
            self.after(0, self.animation_window.stop_animation)
            self.animation_window = None
        self.unbind("<FocusIn>")

    def _lift_animation_window(self, event=None):
        if self.animation_window and self.animation_window.winfo_exists(): self.animation_window.lift()

    def ejecutar_proceso(self, origen='Manual', modo_automatico=False, archivo_salida_auto=None):
        # Validación de ruta previa
        ruta_actual = self.carpeta_entrada.get()
        if not modo_automatico:
            if not ruta_actual or not os.path.isdir(ruta_actual) or not os.listdir(ruta_actual):
                Messagebox.show_warning("La carpeta seleccionada no es válida o está vacía.\nPor favor seleccione una carpeta válida.", "Carpeta Inválida", parent=self)
                self.seleccionar_carpeta()
                if not self.carpeta_entrada.get() or not os.path.isdir(self.carpeta_entrada.get()):
                    return

        if self.piloto_activo and not modo_automatico:
            Messagebox.show_warning("Desactive el Piloto Automático para realizar un proceso manual.", "Piloto Activo", parent=self); return
        threading.Thread(target=self._proceso_thread, args=(origen, modo_automatico, archivo_salida_auto), daemon=True).start()

    def _proceso_thread(self, origen, modo_automatico, archivo_salida_auto):
        self.show_animation("Iniciando Proceso...")

        # Inicialización de variables
        run_id = datetime.now().timestamp()
        if origen == 'Manual':
            self.last_manual_run_id = run_id
        else:
            self.current_auto_run_id = run_id

        periodo, proceso_ok, subida_resumen_ok = "N/A", False, False
        incidencias_resumen, reporte_incidentes = "Error no especificado", pd.DataFrame()
        datos_combinados = pd.DataFrame()
        regional, archivo_salida_final = "Desconocida", None

        try:
            ruta_carpeta = self.carpeta_entrada.get()
            if not ruta_carpeta:
                raise ValueError("Carpeta raíz no seleccionada")

            nombre_carpeta_raiz = os.path.basename(ruta_carpeta)
            regional = obtener_regional(nombre_carpeta_raiz)

            self.after(0, lambda: self.notebook.tab(self.incidentes_tab, text=f" Historial Sinc: {nombre_carpeta_raiz} "))

            if modo_automatico:
                start_date_str = self.date_start.entry.get()
                fecha_inicio_dt = datetime.strptime(start_date_str, '%d/%m/%Y').date()
                fecha_fin_dt = datetime.now().date()
                fecha_fin_query = fecha_fin_dt + timedelta(days=1)
                periodo = f"{start_date_str} - {fecha_fin_dt.strftime('%d/%m/%Y')}"
            else:
                start_date_str, end_date_str = self.date_start.entry.get(), self.date_end.entry.get()
                fecha_inicio_dt = datetime.strptime(start_date_str, '%d/%m/%Y').date()
                fecha_fin_dt = datetime.strptime(end_date_str, '%d/%m/%Y').date()
                fecha_fin_query = fecha_fin_dt + timedelta(days=1)
                periodo = f"{start_date_str} - {end_date_str}"

            datos_combinados, reporte_incidentes = procesar_bases_de_datos(
                ruta_carpeta, password, fecha_inicio_dt, fecha_fin_query, self
            )
            self.resumen = generar_resumen(datos_combinados)
            self.resumen_servicios = generar_resumen_servicios(datos_combinados)
            self.resumen_cruzado = generar_resumen_cruzado(datos_combinados)
            self.reporte_incidentes = reporte_incidentes
            incidencias_resumen = self.summarize_incidents(reporte_incidentes)

            if modo_automatico:
                master_path = os.path.dirname(ruta_carpeta)
                fecha_actual = datetime.now()
                mes_nombre = fecha_actual.strftime('%B')
                output_folder = os.path.join(master_path, "Reportes Generados", regional, mes_nombre)
                os.makedirs(output_folder, exist_ok=True)

                archivo_nombre = f"{regional}_{fecha_inicio_dt.strftime('%d-%m-%Y')}_{fecha_fin_dt.strftime('%d-%m-%Y')}_{fecha_actual.strftime('%H-%M')}.xlsx"
                archivo_salida_final = os.path.join(output_folder, archivo_nombre)
            else:
                self.hide_animation()
                hora_actual_str = datetime.now().strftime('%H-%M')
                nombre_sugerido = f"{regional}_{start_date_str.replace('/', '-')}_{end_date_str.replace('/', '-')}_{hora_actual_str}.xlsx"
                archivo_salida_final = filedialog.asksaveasfilename(
                    initialfile=nombre_sugerido,
                    defaultextension='.xlsx',
                    filetypes=[("Excel files", "*.xlsx")]
                )
                if archivo_salida_final:
                    self.show_animation("Guardando archivo...")

            if archivo_salida_final and not self.resumen.empty:
                if modo_automatico:
                    self.animation_window.update_status("Guardando archivo Excel...")
                proceso_ok = guardar_en_excel(datos_combinados, reporte_incidentes, self.resumen, archivo_salida_final, self.resumen_cruzado)
                
                if proceso_ok:
                    # Crear una versión reducida para el correo
                    archivo_correo = archivo_salida_final.replace('.xlsx', '_Resumen.xlsx')
                    try:
                        with pd.ExcelWriter(archivo_correo, engine='openpyxl') as writer:
                            self.resumen.to_excel(writer, sheet_name='Resumen', index=False)
                            if not self.resumen_servicios.empty:
                                self.resumen_servicios.to_excel(writer, sheet_name='Servicios', index=False)
                            if not getattr(self, 'resumen_cruzado', pd.DataFrame()).empty:
                                self.resumen_cruzado.to_excel(writer, sheet_name='Resumen Tipo usuario y género', index=False)
                    except Exception as e:
                        logging.error(f"Error generando Excel de correo: {e}")
                        archivo_correo = archivo_salida_final

                    if not modo_automatico:
                        self.btn_subir_g.config(state=NORMAL)
                    else:
                        self.animation_window.update_status("Subiendo Resumen a Google Sheets...")
                        subida_resumen_ok = self.subir_resumen_a_sheets(modo_automatico=True)
            elif not modo_automatico:
                self.status_label.config(text="Guardado cancelado o resumen vacío.")

        except Exception as e:
            incidencias_resumen = str(e)
            logging.error(f"Error en _proceso_thread: {e}")

        finally:
            # Enviar correo (una sola vez, asíncrono)
            if origen != 'Prueba':
                if proceso_ok:
                    asunto_correo = f"Estatus de Sincronización - Regional {regional}"
                    cuerpo_html_correo = self.construir_cuerpo_correo(reporte_incidentes, regional)
                    destinatarios_filtrados = [r['correo'] for r in self.recipients_list if r['regional'] in [regional, "Todas"] and r.get('activo', True)]
                    threading.Thread(
                        target=enviar_correo_notificacion,
                        args=(asunto_correo, cuerpo_html_correo, destinatarios_filtrados, locals().get('archivo_correo', archivo_salida_final)),
                        daemon=True
                    ).start()
                    correo_enviado_ok = True
                else:
                    asunto_correo = f"FALLO en Ejecución del Analizador - {regional}"
                    cuerpo_html_correo = self.construir_cuerpo_correo_de_fallo(regional, periodo, origen, incidencias_resumen)
                    destinatario_admin = [ADMIN_EMAIL]
                    threading.Thread(
                        target=enviar_correo_notificacion,
                        args=(asunto_correo, cuerpo_html_correo, destinatario_admin),
                        daemon=True
                    ).start()
                    correo_enviado_ok = True
            else:
                correo_enviado_ok = False

            # Registrar en historial
            entry_id = run_id
            subida_str = f"{'Sí' if subida_resumen_ok else 'No'}/No"
            history_entry = {
                "id": entry_id,
                "fecha_hora": datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
                "origen": origen,
                "regional": regional,
                "periodo": periodo,
                "archivos": len(reporte_incidentes),
                "registros": len(datos_combinados),
                "proceso": "Sí" if proceso_ok else "No",
                "subida": subida_str,
                "correo_enviado": "Sí" if correo_enviado_ok else "No",
                "incidencias": incidencias_resumen
            }
            self.after(0, lambda: self.add_history_entry(history_entry))
            self.after(0, lambda: self.update_incidents_table(reporte_incidentes))

            # --- REGISTRAR EN HISTORIAL (NUEVO) ---
            if not reporte_incidentes.empty:
                 self.after(0, lambda: sucursales_history.registrar_historial(reporte_incidentes, regional))


            # Subida del historial de sincronización
            if proceso_ok:
                threading.Thread(
                    target=subir_historial_sincronizacion,
                    args=(reporte_incidentes, os.path.basename(self.carpeta_entrada.get())),
                    daemon=True
                ).start()
                
                # Subida de porcentajes de sincronización
                threading.Thread(
                    target=subir_porcentajes_sincronizacion,
                    args=(reporte_incidentes, regional),
                    daemon=True
                ).start()

            self.hide_animation()
    
    

    def extraer_numero_nombre(self, folder_name):
        import re
        match = re.match(r'^(\d+)\s*-\s*(.+)$', str(folder_name))
        if match:
            return int(match.group(1)), match.group(2).strip()
        return None, str(folder_name)

    def subir_a_supabase(self, df_servicios, df_demografico, regional, reporte_incidentes=None, df_cruzado=None):
        if not getattr(self, 'supabase_url', None) or not getattr(self, 'supabase_key', None):
            logging.warning("Credenciales de Supabase no configuradas. Saltando sincronización.")
            return False
            
        import requests
        import json
        import re
        import glob
        
        regional_normalizada = "Chiriquí" if regional == "Chiriqui" else regional
        
        # 1. Escanear carpetas físicas para sincronizar la tabla infoplazas
        ruta_carpeta = self.carpeta_entrada.get()
        if not ruta_carpeta or not os.path.exists(ruta_carpeta):
            logging.warning(f"Ruta de carpeta no válida para Supabase: {ruta_carpeta}")
            return False
            
        try:
            subcarpetas = [f.path for f in os.scandir(ruta_carpeta) if f.is_dir()]
            infoplazas_payload = []
            for carpeta in subcarpetas:
                folder_name = os.path.basename(carpeta)
                num, name = self.extraer_numero_nombre(folder_name)
                if num is None:
                    continue
                
                estado = "Activa"
                if glob.glob(os.path.join(carpeta, "CERRADA DEFINITIVAMENTE.txt")):
                    estado = "Cerrada Definitivamente"
                    
                infoplazas_payload.append({
                    "numero": num,
                    "nombre": name,
                    "nombre_carpeta": folder_name,
                    "regional": regional_normalizada,
                    "estado": estado
                })
                
            headers = {
                "apikey": self.supabase_key,
                "Authorization": f"Bearer {self.supabase_key}",
                "Content-Type": "application/json"
            }
            
            # Subir infoplazas
            if infoplazas_payload:
                url_rpc_info = f"{self.supabase_url.rstrip('/')}/rest/v1/rpc/upsert_infoplazas"
                res = requests.post(url_rpc_info, headers=headers, json={"payload": infoplazas_payload})
                if res.status_code not in [200, 201, 204]:
                    logging.error(f"Error al subir infoplazas a Supabase: {res.status_code} - {res.text}")
                else:
                    logging.info(f"Sincronizadas {len(infoplazas_payload)} infoplazas en Supabase.")
                    
            # 2. Subir resumen de servicios
            servicios_payload = []
            servicios_map = {
                'TALLER': 'taller', 'REUNIÓN': 'reunion', 'CONSULTA': 'consulta',
                'VENTA': 'venta', 'SCAN': 'scan', 'CORREO': 'correo',
                'TEL': 'tel', 'LT': 'lt', 'IMPRESIÓN': 'impresion',
                'COPIA': 'copia', 'CINE': 'cine', 'OTROS': 'otros',
                'USO DE PC': 'uso_de_pc'
            }
            
            if df_servicios is not None and not df_servicios.empty:
                for _, row in df_servicios.iterrows():
                    infoplaza_str = str(row['Infoplaza'])
                    num, _ = self.extraer_numero_nombre(infoplaza_str)
                    if num is None:
                        continue
                    record = {
                        'numero_infoplaza': num,
                        'regional': regional_normalizada,
                        'anio': int(row['Año']),
                        'mes': str(row['Mes']).strip(),
                        'total': int(row['Total'])
                    }
                    for py_col, db_col in servicios_map.items():
                        record[db_col] = int(row.get(py_col, 0))
                    servicios_payload.append(record)
                    
            if servicios_payload:
                url_rpc_servicios = f"{self.supabase_url.rstrip('/')}/rest/v1/rpc/upsert_resumen_servicios"
                res = requests.post(url_rpc_servicios, headers=headers, json={"payload": servicios_payload})
                if res.status_code not in [200, 201, 204]:
                    logging.error(f"Error al subir resumen_servicios a Supabase: {res.status_code} - {res.text}")
                else:
                    logging.info(f"Sincronizado resumen_servicios en Supabase ({len(servicios_payload)} filas).")
                    
            # 3. Subir resumen demográfico
            demograficos_payload = []
            demograficos_map = {
                'Masculino': 'masculino', 'Femenino': 'femenino',
                'Primaria': 'primaria', 'Secundaria': 'secundaria',
                'Universitario': 'universitario', 'Docente': 'docente',
                'Tercera Edad': 'tercera_edad', 'Público General': 'publico_general'
            }
            
            if df_demografico is not None and not df_demografico.empty:
                for _, row in df_demografico.iterrows():
                    infoplaza_str = str(row['Infoplaza'])
                    num, _ = self.extraer_numero_nombre(infoplaza_str)
                    if num is None:
                        continue
                    record = {
                        'numero_infoplaza': num,
                        'regional': regional_normalizada,
                        'anio': int(row['Año']),
                        'mes': str(row['Mes']).strip(),
                        'total': int(row['Total'])
                    }
                    for py_col, db_col in demograficos_map.items():
                        record[db_col] = int(row.get(py_col, 0))
                    demograficos_payload.append(record)
                    
            if demograficos_payload:
                url_rpc_demo = f"{self.supabase_url.rstrip('/')}/rest/v1/rpc/upsert_resumen_demografico"
                res = requests.post(url_rpc_demo, headers=headers, json={"payload": demograficos_payload})
                if res.status_code not in [200, 201, 204]:
                    logging.error(f"Error al subir resumen_demografico a Supabase: {res.status_code} - {res.text}")
                else:
                    logging.info(f"Sincronizado resumen_demografico en Supabase ({len(demograficos_payload)} filas).")
                    
            # 3.5 Subir resumen tipo usuario y genero (cruzado)
            cruzado_payload = []
            
            if df_cruzado is not None and not df_cruzado.empty:
                for _, row in df_cruzado.iterrows():
                    infoplaza_str = str(row['Infoplaza'])
                    num, _ = self.extraer_numero_nombre(infoplaza_str)
                    if num is None:
                        continue
                    record = {
                        'numero_infoplaza': num,
                        'regional': regional_normalizada,
                        'anio': int(row['Año']),
                        'mes': str(row['Mes']).strip(),
                        'tipo_usuario': str(row['Tipo Usuario']),
                        'masculino': int(row.get('Masculino', 0)),
                        'femenino': int(row.get('Femenino', 0)),
                        'total': int(row['Total'])
                    }
                    cruzado_payload.append(record)
                    
            if cruzado_payload:
                url_rpc_cruzado = f"{self.supabase_url.rstrip('/')}/rest/v1/rpc/upsert_resumen_tipo_usuario_genero"
                res = requests.post(url_rpc_cruzado, headers=headers, json={"payload": cruzado_payload})
                if res.status_code not in [200, 201, 204]:
                    logging.error(f"Error al subir resumen_tipo_usuario_genero a Supabase: {res.status_code} - {res.text}")
                else:
                    logging.info(f"Sincronizado resumen_tipo_usuario_genero en Supabase ({len(cruzado_payload)} filas).")
                    
            # 4. Subir historial de sincronización
            if reporte_incidentes is not None and not reporte_incidentes.empty:
                historial_payload = []
                for _, row in reporte_incidentes.iterrows():
                    sucursal_str = str(row['Carpeta'])
                    dias_val = row.get('Días sin Sincronizar', 0)
                    if dias_val == 'N/A' or pd.isna(dias_val):
                        dias_val = 0
                    else:
                        try:
                            dias_val = int(dias_val)
                        except:
                            dias_val = 0
                            
                    obs_val = str(row.get('Observación', '')).strip()
                    
                    historial_payload.append({
                        'fecha_reporte': datetime.now().strftime('%Y-%m-%d'),
                        'regional': regional_normalizada,
                        'sucursal': sucursal_str,
                        'dias_sin_sinc': dias_val,
                        'observacion': obs_val
                    })
                    
                if historial_payload:
                    url_historial = f"{self.supabase_url.rstrip('/')}/rest/v1/historial_sincronizacion"
                    res = requests.post(url_historial, headers=headers, json=historial_payload)
                    if res.status_code not in [200, 201, 204]:
                        logging.error(f"Error al subir historial_sincronizacion a Supabase: {res.status_code} - {res.text}")
                    else:
                        logging.info(f"Sincronizado historial_sincronizacion en Supabase ({len(historial_payload)} filas).")
            
            return True
        except Exception as e:
            logging.error(f"Excepción en subir_a_supabase: {e}")
            return False

    def subir_resumen_a_sheets(self, modo_automatico=False):
        if self.resumen.empty:
            if not modo_automatico: Messagebox.show_warning("No hay datos de resumen para subir.", "Datos Vacíos", parent=self)
            return False
        
        if not modo_automatico: self.show_animation("Subiendo a Google Sheets...")
        subida_ok = False
        try:
            nombre_carpeta = os.path.basename(self.carpeta_entrada.get())
            subida_ok = subir_datos_a_google_sheets(self.resumen, nombre_carpeta)
            
            # --- SUBIDA EN PARALELO A SUPABASE (ASÍNCRONA VÍA HILO) ---
            def run_supabase_upload():
                supabase_ok = False
                try:
                    regional = obtener_regional(nombre_carpeta)
                    supabase_ok = self.subir_a_supabase(
                        getattr(self, 'resumen_servicios', pd.DataFrame()),
                        self.resumen,
                        regional,
                        getattr(self, 'reporte_incidentes', None),
                        getattr(self, 'resumen_cruzado', pd.DataFrame())
                    )
                except Exception as se:
                    logging.error(f"Fallo al sincronizar con Supabase: {se}")
                
                # Actualizar el estatus en el historial con el callback combinado
                entry_id = getattr(self, 'current_auto_run_id', None) if modo_automatico else self.last_manual_run_id
                if entry_id is not None:
                    self.actualizar_estado_subida_historial(entry_id, subida_ok, supabase_ok)

            threading.Thread(target=run_supabase_upload, daemon=True).start()

            if subida_ok and not modo_automatico:
                Messagebox.show_info(f"Los datos se han subido correctamente a la hoja '{nombre_carpeta}'.", "Subida Exitosa", parent=self)
                self.status_label.config(text="Datos subidos a la nube.")
            return subida_ok
        except Exception as e:
            logging.error(f"FALLO en subida a Google Sheets: {e}")
            if not modo_automatico: Messagebox.show_error(f"No se pudo subir a Google Sheets:\n{e}", "Error de Subida", parent=self)
            return False
        finally:
            if not modo_automatico: self.hide_animation()



    def construir_cuerpo_correo(self, reporte_df, regional):
        """Genera el cuerpo HTML del correo con formato solicitado y compatible con Outlook."""

        if reporte_df.empty:
            return "<p>No se generó reporte de incidentes para esta ejecución.</p>"

        # --- Separar cerradas de operativas ---
        df_cerradas = reporte_df[reporte_df['Observación'] == 'CERRADA DEFINITIVAMENTE'].copy()
        df_operativas = reporte_df[reporte_df['Observación'] != 'CERRADA DEFINITIVAMENTE'].copy()

        # Asegurar columna numérica
        df_operativas['Días sin Sincronizar'] = pd.to_numeric(df_operativas['Días sin Sincronizar'], errors='coerce').fillna(0)

        # --- Calcular métricas para la ejecución ACTUAL (Hoy) ---
        total_verificadas = len(df_operativas)
        al_dia = len(df_operativas[df_operativas['Días sin Sincronizar'] <= 10])
        
        # Para revisión: Unión para evitar doble conteo (Especialmente en carpetas vacías con muchos días)
        df_para_revision_hoy = df_operativas[
            (df_operativas['Observación'] == 'Carpeta Vacía') | 
            (df_operativas['Días sin Sincronizar'] > 10)
        ]
        para_revision = len(df_para_revision_hoy)
        
        # Calcular porcentajes para hoy
        porc_al_dia_hoy = round((al_dia / total_verificadas) * 100, 1) if total_verificadas > 0 else 0
        porc_revision_hoy = round((para_revision / total_verificadas) * 100, 1) if total_verificadas > 0 else 0

        # --- Obtener PROMEDIOS MENSUALES desde Google Sheets e incorporar HOY ---
        prom_hist, prom_rev_hist, n_hist = obtener_promedios_mensuales_sheets(regional)
        
        if prom_hist is not None:
            # Recalcular promedio incluyendo la ejecución de hoy para mayor precisión
            n_ejec = n_hist + 1
            prom_al_dia = round(((prom_hist * n_hist) + porc_al_dia_hoy) / n_ejec, 1)
            prom_rev = round(((prom_rev_hist * n_hist) + porc_revision_hoy) / n_ejec, 1)
        else:
            # Primera ejecución del mes
            prom_al_dia, prom_rev, n_ejec = porc_al_dia_hoy, porc_revision_hoy, 1

        # --- Saludo dinámico ---
        saludo = "Buen día" if datetime.now().hour < 12 else "Buenas tardes"

        # --- Datos del reporte ---
        fecha_reporte = datetime.now().strftime('%d de %B de %Y')
        mes_reporte = datetime.now().strftime('%B de %Y')

        # --- Helper mejorado para tablas con diseño moderno ---
        def crear_tabla_html(titulo, df, columnas, color_header, emoji, agregar_meses=False, resaltar_dias_criticos=False):
            if df.empty:
                return ""
            
            df_tabla = df.copy()
            df_tabla.insert(0, '#', range(1, 1 + len(df_tabla)))
            df_tabla.rename(columns={
                'Carpeta': 'Sucursal',
                'Fecha Archivo': 'Última sincronización',
                'Fecha y Hora': 'Fecha'
            }, inplace=True)

            dias_originales = None
            if 'Días sin Sincronizar' in df_tabla.columns:
                dias_originales = pd.to_numeric(df_tabla['Días sin Sincronizar'], errors='coerce').fillna(0).astype(int)
                df_tabla['Días sin Sincronizar'] = dias_originales
                if agregar_meses:
                    df_tabla['Días sin Sincronizar'] = df_tabla['Días sin Sincronizar'].apply(
                        lambda d: f"{d}{_calcular_meses_relativos(d)}" if d > 30 else str(d)
                    )

            columnas_existentes = ['#'] + [col for col in columnas if col in df_tabla.columns]
            html_tabla = df_tabla[columnas_existentes].to_html(index=False, border=0, na_rep='', escape=False)
            
            if resaltar_dias_criticos and dias_originales is not None:
                filas_html = html_tabla.split('<tr>')
                nueva_html = [filas_html[0]]
                for i, fila in enumerate(filas_html[1:], start=0):
                    if i < len(dias_originales):
                        dias = dias_originales.iloc[i]
                        if dias == 10: fila = fila.replace('<tr>', '<tr style="background-color: #FFE4CC;">', 1)
                        elif dias == 9: fila = fila.replace('<tr>', '<tr style="background-color: #FFF4CC;">', 1)
                    nueva_html.append(fila)
                html_tabla = '<tr>'.join(nueva_html)
            
            count_badge = f'<span style="background: rgba(255,255,255,0.3); padding: 4px 12px; border-radius: 20px; font-size: 14px; font-weight: 700; margin-left: 10px;">{len(df_tabla)}</span>'
            
            return f"""
            <div style="margin: 25px 0; border-radius: 10px; overflow: hidden; box-shadow: 0 3px 8px rgba(0,0,0,0.08);">
                <div style="background: {color_header}; color: white; padding: 15px 20px; font-size: 16px; font-weight: 600;">
                    {emoji} {titulo} {count_badge}
                </div>
                {html_tabla}
            </div>
            """

        # --- Clasificación de tablas ---
        df_carpetas_vacias = df_operativas[df_operativas['Observación'] == 'Carpeta Vacía'].copy()
        df_trabajo = df_operativas[df_operativas['Observación'] != 'Carpeta Vacía'].copy()
        observaciones_basicas = ['OK', 'Sin registros para el periodo', 'CERRADA DEFINITIVAMENTE', 'CERRADA TEMPORALMENTE']
        
        df_con_observaciones = df_trabajo[(df_trabajo['Días sin Sincronizar'] > 10) & (~df_trabajo['Observación'].isin(observaciones_basicas))].copy().sort_values(by='Días sin Sincronizar', ascending=False)
        indices_rojas = df_con_observaciones.index.tolist()
        df_revision_dias = df_trabajo[(~df_trabajo.index.isin(indices_rojas)) & (df_trabajo['Días sin Sincronizar'] > 10) & (df_trabajo['Observación'].isin(['OK', 'Sin registros para el periodo']))].copy().sort_values(by='Días sin Sincronizar', ascending=False)
        if not df_revision_dias.empty: df_revision_dias['Observación'] = 'Revisar'
        
        indices_todas_revision = indices_rojas + df_revision_dias.index.tolist()
        df_en_periodo = df_trabajo[(df_trabajo['Días sin Sincronizar'] <= 10)].copy().sort_values(by='Días sin Sincronizar', ascending=False)

        html_carpetas_vacias = crear_tabla_html('Sucursales con Carpeta Vacía', df_carpetas_vacias, ['Sucursal'], '#e74c3c', '🚩') if not df_carpetas_vacias.empty else ""
        html_observaciones = crear_tabla_html('Sucursales con Observaciones', df_con_observaciones, ['Sucursal', 'Última sincronización', 'Días sin Sincronizar', 'Observación'], '#c0392b', '🔴', True)
        html_revision = crear_tabla_html('Sucursales con Más de 10 Días sin Sincronizar', df_revision_dias, ['Sucursal', 'Última sincronización', 'Días sin Sincronizar', 'Observación'], '#f39c12', '🟡', True)
        html_sincronizadas = crear_tabla_html('Sucursales Sincronizadas Dentro del Periodo', df_en_periodo, ['Sucursal', 'Última sincronización', 'Días sin Sincronizar'], '#27ae60', '✅', False, True)
        html_cerradas = crear_tabla_html('Sucursales Cerradas Definitivamente', df_cerradas, ['Sucursal', 'Observación'], '#7f8c8d', '🔒')

        ameritan_revision_count = len(df_con_observaciones) + len(df_revision_dias)

        # --- Ensamblaje del correo ---
        cuerpo_completo = f"""
        <html>
        <head>
            <meta http-equiv="Content-Type" content="text/html; charset=utf-8">
            <meta name="viewport" content="width=device-width, initial-scale=1.0">
            <style type="text/css">
                body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, Arial, sans-serif; font-size: 14px; line-height: 1.6; color: #333; margin: 0; padding: 20px; background-color: #f5f7fa; }}
                .email-container {{ max-width: 900px; margin: 0 auto; background: white; border-radius: 12px; overflow: hidden; box-shadow: 0 8px 16px rgba(0,0,0,0.1); }}
                .header {{ background-color: #667eea; color: white; padding: 30px; text-align: left; }}
                .content {{ padding: 30px; }}
                table {{ border-collapse: collapse; width: 100%; background: white; }}
                th, td {{ text-align: left; padding: 12px 15px; border-bottom: 1px solid #f0f0f0; }}
                th {{ background-color: #f8f9fa; font-weight: 600; color: #495057; font-size: 13px; text-transform: uppercase; }}
                .info-box {{ background-color: #e3f2fd; border-left: 4px solid #2196f3; border-radius: 8px; padding: 20px; margin: 20px 0; }}
                .porc-sinc-container {{ background-color: #f5f7fa; border-radius: 12px; padding: 25px; margin: 25px 0; box-shadow: 0 4px 12px rgba(0,0,0,0.08); }}
                .porc-sinc-title {{ font-size: 20px; font-weight: 700; color: #2c3e50; margin-bottom: 20px; text-align: center; padding-bottom: 15px; border-bottom: 3px solid #667eea; }}
                .stat-card {{ background-color: white; border-radius: 10px; padding: 20px; margin: 10px 0; box-shadow: 0 3px 8px rgba(0,0,0,0.08); border-top: 4px solid #667eea; text-align: center; }}
                .stat-value {{ font-size: 32px; font-weight: 700; color: #2c3e50; margin: 5px 0; }}
                .stat-label {{ font-size: 12px; color: #7f8c8d; text-transform: uppercase; letter-spacing: 0.5px; font-weight: 600; }}
                .exec-today-pct {{ font-size: 13px; color: #7f8c8d; background: #f0f2f5; display: inline-block; padding: 2px 10px; border-radius: 12px; margin-top: 8px; font-weight: 600; border: 1px solid #e1e4e8; }}
                .info-note {{ background-color: #e8f5e9; border-left: 4px solid #4caf50; padding: 15px 20px; border-radius: 8px; margin: 20px 0; font-size: 13px; color: #1b5e20; }}
                .attachment-note {{ background-color: #fff3e0; border-left: 4px solid #ff9800; padding: 12px 20px; border-radius: 8px; margin: 20px 0; font-size: 14px; color: #e65100; }}
                .footer {{ margin-top: 40px; padding: 25px 30px; background-color: #f8f9fa; border-top: 3px solid #667eea; }}
            </style>
        </head>
        <body>
            <div class="email-container">
                <div class="header">
                    <h2 style="margin: 0; font-size: 24px; color: white; font-weight: 700;">📊 Reporte de Sincronización - OneRoof</h2>
                    <p style="margin: 8px 0 0 0; font-size: 14px; color: white;">Sistema ADA | Analizador de Bases de Datos</p>
                </div>
                
                <div class="content">
                    <p style="font-size: 16px; margin-bottom: 20px; color: #2c3e50;"><strong>{saludo},</strong></p>
                    <p>Le compartimos el estatus actualizado de la sincronización regional de la <strong>Regional de {regional}</strong>.</p>
                    
                    <div class="attachment-note">📎 <strong>Archivo adjunto:</strong> Se incluye documento en Excel detallado.</div>
                    
                    <div class="info-box">
                        <h3 style="margin: 0 0 10px 0; font-size: 16px; color: #1565c0;">📋 Datos del Reporte</h3>
                        <ul style="margin: 0; padding-left: 20px;">
                            <li><strong>Regional:</strong> {regional}</li>
                            <li><strong>Fecha:</strong> {fecha_reporte}</li>
                        </ul>
                    </div>
                    
                    <div class="porc-sinc-container">
                        <div class="porc-sinc-title">📈 Promedio Mensual de Sincronización - {mes_reporte}</div>
                        <table width="100%" cellpadding="0" cellspacing="0" border="0">
                            <tr>
                                <td width="48%" valign="top">
                                    <div class="stat-card">
                                        <div class="stat-label">✅ PROMEDIO AL DÍA</div>
                                        <div class="stat-value">{prom_al_dia}%</div>
                                        <div style="width: 100%; height: 10px; background-color: #ecf0f1; border-radius: 5px; overflow: hidden; margin-top: 10px;">
                                            <div style="height: 100%; width: {prom_al_dia}%; background-color: #667eea; border-radius: 5px;"></div>
                                        </div>
                                        <div style="font-size: 11px; color: #bdc3c7; margin-top: 8px;">(Basado en {n_ejec} ejecuciones este mes)</div>
                                    </div>
                                </td>
                                <td width="4%"></td>
                                <td width="48%" valign="top">
                                    <div class="stat-card" style="border-top-color: #f39c12;">
                                        <div class="stat-label">⚠️ PROMEDIO PARA REVISIÓN</div>
                                        <div class="stat-value">{prom_rev}%</div>
                                        <div style="width: 100%; height: 10px; background-color: #ecf0f1; border-radius: 5px; overflow: hidden; margin-top: 10px;">
                                            <div style="height: 100%; width: {prom_rev}%; background-color: #f39c12; border-radius: 5px;"></div>
                                        </div>
                                        <div style="font-size: 11px; color: #bdc3c7; margin-top: 8px;">(Basado en {n_ejec} ejecuciones este mes)</div>
                                    </div>
                                </td>
                            </tr>
                        </table>
                    </div>

                    <div style="margin: 30px 0;">
                        <div style="font-size: 20px; font-weight: 700; color: #2c3e50; margin-bottom: 20px; padding-bottom: 10px; border-bottom: 3px solid #667eea;">📊 Resumen Ejecutivo (Hoy)</div>
                        <table width="100%" cellpadding="0" cellspacing="0" border="0">
                            <tr>
                                <td width="31%" valign="top">
                                    <div class="stat-card">
                                        <div style="font-size: 28px; margin-bottom: 10px;">📈</div>
                                        <div class="stat-value">{total_verificadas}</div>
                                        <div class="stat-label">Verificadas</div>
                                    </div>
                                </td>
                                <td width="3.5%"></td>
                                <td width="31%" valign="top">
                                    <div class="stat-card">
                                        <div style="font-size: 28px; margin-bottom: 10px;">✅</div>
                                        <div class="stat-value">{al_dia}</div>
                                        <div class="stat-label">En Periodo</div>
                                        <div class="exec-today-pct">{porc_al_dia_hoy}%</div>
                                    </div>
                                </td>
                                <td width="3.5%"></td>
                                <td width="31%" valign="top">
                                    <div class="stat-card" style="border-top-color: #f39c12;">
                                        <div style="font-size: 28px; margin-bottom: 10px;">⚠️</div>
                                        <div class="stat-value">{ameritan_revision_count}</div>
                                        <div class="stat-label">Para Revisión</div>
                                        <div class="exec-today-pct">{porc_revision_hoy}%</div>
                                    </div>
                                </td>
                            </tr>
                        </table>
                    </div>
                    
                    <div class="info-note"><strong>Nota informativa:</strong> Los indicadores superiores reflejan la tendencia mensual acumulada, mientras que el resumen ejecutivo muestra el estado exacto de la última ejecución.</div>
                    
                    <hr>
                    {html_carpetas_vacias}
                    {html_observaciones}
                    {html_revision}
                    {html_sincronizadas}
                    {html_cerradas}
                </div>
                
                <div class="footer">
                    <p style="margin: 8px 0; color: #495057;">Atentamente,</p>
                    <p style="font-size: 18px; font-weight: 700; color: #667eea; margin-top: 10px;">Equipo ADA | GEBILO Regional de Los Santos</p>
                    <div style="font-size: 11px; color: #868e96; font-style: italic; padding-top: 15px; border-top: 1px solid #dee2e6;">Este es un mensaje automático del Sistema ADA. Confidencial y de uso exclusivo.</div>
                </div>
            </div>
        </body>
        </html>
        """
        return cuerpo_completo
    

    def iniciar_piloto(self, from_load=False, proxima_ejecucion_guardada=None):
        # Validación de ruta previa
        ruta_actual = self.carpeta_entrada.get()
        if not ruta_actual or not os.path.isdir(ruta_actual) or not os.listdir(ruta_actual):
            if not from_load: # Solo mostrar advertencia si es acción manual del usuario, no al cargar app
                Messagebox.show_warning("La carpeta seleccionada no es válida o está vacía.\nPor favor seleccione una carpeta válida.", "Carpeta Inválida", parent=self)
                self.seleccionar_carpeta()
            
            # Verificar de nuevo
            ruta_nueva = self.carpeta_entrada.get()
            if not ruta_nueva or not os.path.isdir(ruta_nueva) or not os.listdir(ruta_nueva):
                 if not from_load: Messagebox.show_error("No se puede iniciar el piloto sin una carpeta válida.", "Error", parent=self)
                 return

        # Obtiene la lista de regionales seleccionadas desde los checkboxes
        regiones_seleccionadas = [region for region, var in self.regional_vars.items() if var.get()]
        
        if not regiones_seleccionadas:
            Messagebox.show_warning("Seleccione al menos una regional para el Piloto Automático.", "Configuración Requerida", parent=self)
            return

        dias_map = {"L": "Lunes", "M": "Martes", "X": "Miércoles", "J": "Jueves", "V": "Viernes", "S": "Sábado", "D": "Domingo"}
        dias = [dias_map[dia_key] for dia_key, var in self.vars_dias.items() if var.get()]
        if not dias:
            Messagebox.show_warning("Selecciona al menos un día para la ejecución semanal.", "Configuración Requerida", parent=self); return
        
        self.piloto_activo = True; self.btn_iniciar_piloto.config(text="Detener", bootstyle="danger", command=self.detener_piloto)
        self.piloto_switch.config(state=DISABLED); self.status_label.config(text="Piloto Automático ACTIVO.")
        
        # Pasa la lista de regionales al hilo del piloto
        self.hilo_piloto = threading.Thread(target=bucle_piloto_automatico, args=(self, dias, regiones_seleccionadas, proxima_ejecucion_guardada, not from_load), daemon=True)
        self.hilo_piloto.start()
        self.save_settings()

    def detener_piloto(self):
        self.piloto_activo = False; self.btn_iniciar_piloto.config(text="Iniciar", bootstyle="success", command=self.iniciar_piloto)
        self.piloto_switch.config(state=NORMAL); self.status_label.config(text="Piloto Automático detenido.")
        self.proxima_ejecucion_var.set(""); self.proxima_ejecucion_dt = None; self.save_settings()

    def actualizar_cronometro(self, proxima_ejecucion_dt):
        if not self.piloto_activo: self.proxima_ejecucion_var.set(""); return
        segundos_restantes = (proxima_ejecucion_dt - datetime.now()).total_seconds()
        if segundos_restantes > 0:
            dias, rem = divmod(segundos_restantes, 86400); horas, rem = divmod(rem, 3600); minutos, segundos = divmod(rem, 60)
            fecha_str = proxima_ejecucion_dt.strftime('%d/%m/%Y a las %H:%M'); countdown_str = f"{int(dias)}d {int(horas):02}h {int(minutos):02}m {int(segundos):02}s"
            self.proxima_ejecucion_var.set(f"Próxima ejecución: {fecha_str} (en {countdown_str})")
            self.after(1000, lambda: self.actualizar_cronometro(proxima_ejecucion_dt))
        else: self.proxima_ejecucion_var.set("Ejecutando ahora...")
    
    def summarize_incidents(self, reporte_df):
        if reporte_df.empty or all(obs == "OK" for obs in reporte_df['Observación']): return "OK"
        errores = reporte_df[reporte_df['Observación'].str.contains("Error", na=False)]; sin_datos = reporte_df[reporte_df['Observación'].str.contains("Sin datos|Fuera de rango|CERRADA|Sin archivo|Carpeta Vacía", na=False)]
        resumen = []
        if not errores.empty: resumen.append(f"{len(errores)} con error")
        if not sin_datos.empty: resumen.append(f"{len(sin_datos)} con obs.")
        return ", ".join(resumen) if resumen else "OK"

    def add_history_entry(self, entry):
        self.execution_history.insert(0, entry);
        if len(self.execution_history) > 100: self.execution_history.pop()
        self.update_history_table(); self.save_settings()
        
        def run_upload_db():
            if not getattr(self, 'supabase_url', None) or not getattr(self, 'supabase_key', None):
                return
            try:
                import requests
                url = f"{self.supabase_url.rstrip('/')}/rest/v1/historial_ejecuciones"
                headers = {
                    "apikey": self.supabase_key,
                    "Authorization": f"Bearer {self.supabase_key}",
                    "Content-Type": "application/json",
                    "Prefer": "resolution=merge-duplicates"
                }
                payload = {
                    "id": entry.get('id'),
                    "fecha_hora": entry.get('fecha_hora', ''),
                    "origen": entry.get('origen', ''),
                    "regional": entry.get('regional', ''),
                    "periodo": entry.get('periodo', ''),
                    "archivos": int(entry.get('archivos', 0)),
                    "registros": int(entry.get('registros', 0)),
                    "proceso": entry.get('proceso', 'No'),
                    "subida": entry.get('subida', 'No/No'),
                    "correo_enviado": entry.get('correo_enviado', 'No'),
                    "incidencias": entry.get('incidencias', '')
                }
                res = requests.post(url, headers=headers, json=payload)
                if res.status_code not in [200, 201, 204]:
                    logging.error(f"Error al subir historial_ejecuciones a Supabase: {res.status_code} - {res.text}")
            except Exception as ex:
                logging.error(f"Fallo al subir historial de ejecución a Supabase: {ex}")
                
        threading.Thread(target=run_upload_db, daemon=True).start()

    def update_history_table(self):
        for item in self.history_tree.get_children(): self.history_tree.delete(item)
        for i, entry in enumerate(self.execution_history):
            tag = 'oddrow' if i % 2 == 0 else 'evenrow'
            subida_val = entry.get('subida', 'No')
            if subida_val == 'Sí':
                subida_val = 'Sí/No'
            elif subida_val == 'No':
                subida_val = 'No/No'
            
            values = (i + 1, entry.get('fecha_hora', ''), entry.get('origen', ''), entry.get('regional', ''), entry.get('periodo', ''), entry.get('archivos', ''), entry.get('registros', ''), entry.get('proceso', 'No'), subida_val, entry.get('correo_enviado', 'No'), entry.get('incidencias', ''))
            self.history_tree.insert('', END, values=values, tags=(tag,))

    def actualizar_estado_subida_historial(self, entry_id, sheets_ok, supabase_ok):
        target_entry = None
        for entry in self.execution_history:
            if entry.get('id') == entry_id:
                sheets_str = "Sí" if sheets_ok else "No"
                supabase_str = "Sí" if supabase_ok else "No"
                entry['subida'] = f"{sheets_str}/{supabase_str}"
                target_entry = entry
                break
        self.save_settings()
        self.after(0, self.update_history_table)
        
        if target_entry:
            def run_update_db():
                if not getattr(self, 'supabase_url', None) or not getattr(self, 'supabase_key', None):
                    return
                try:
                    import requests
                    url = f"{self.supabase_url.rstrip('/')}/rest/v1/historial_ejecuciones"
                    headers = {
                        "apikey": self.supabase_key,
                        "Authorization": f"Bearer {self.supabase_key}",
                        "Content-Type": "application/json",
                        "Prefer": "resolution=merge-duplicates"
                    }
                    payload = {
                        "id": target_entry.get('id'),
                        "fecha_hora": target_entry.get('fecha_hora', ''),
                        "origen": target_entry.get('origen', ''),
                        "regional": target_entry.get('regional', ''),
                        "periodo": target_entry.get('periodo', ''),
                        "archivos": int(target_entry.get('archivos', 0)),
                        "registros": int(target_entry.get('registros', 0)),
                        "proceso": target_entry.get('proceso', 'No'),
                        "subida": target_entry.get('subida', 'No/No'),
                        "correo_enviado": target_entry.get('correo_enviado', 'No'),
                        "incidencias": target_entry.get('incidencias', '')
                    }
                    res = requests.post(url, headers=headers, json=payload)
                    if res.status_code not in [200, 201, 204]:
                        logging.error(f"Error al actualizar historial_ejecuciones a Supabase: {res.status_code} - {res.text}")
                except Exception as ex:
                    logging.error(f"Fallo al actualizar historial de ejecución en Supabase: {ex}")
                    
            threading.Thread(target=run_update_db, daemon=True).start()

    def enviar_correo_consolidado_automatico(self):
        if not getattr(self, 'resultados_piloto_actual', None):
            return
        
        logging.info("PILOTO AUTOMÁTICO: Generando y enviando correo consolidado...")
        fecha_actual = datetime.now().strftime('%d de %B de %Y')
        asunto = f"Reporte Consolidado de Sincronización - Piloto Automático - {fecha_actual}"
        
        # Recopilar destinatarios
        regiones_ejecutadas = [r['regional'] for r in self.resultados_piloto_actual]
        destinatarios_set = set()
        for rec in self.recipients_list:
            if rec.get('activo', True):
                if rec['regional'] == "Todas" or rec['regional'] in regiones_ejecutadas:
                    destinatarios_set.add(rec['correo'])
        
        destinatarios = list(destinatarios_set)
        if not destinatarios:
            destinatarios = [ADMIN_EMAIL]
            logging.warning(f"No hay destinatarios activos para las regionales ejecutadas. Enviando a respaldo admin: {ADMIN_EMAIL}")

        # Construir tabla resumen general
        filas_resumen = []
        for r in self.resultados_piloto_actual:
            estado_badge = '<span class="badge-exito" style="background-color: #c6f6d5; color: #22543d; padding: 4px 8px; border-radius: 4px; font-weight: bold; font-size: 11px; display: inline-block;">ÉXITO</span>' if r['proceso_ok'] else '<span class="badge-fallo" style="background-color: #fed7d7; color: #742a2a; padding: 4px 8px; border-radius: 4px; font-weight: bold; font-size: 11px; display: inline-block;">FALLO</span>'
            excel_status = "Excel Adjuntado" if (r['proceso_ok'] and r['archivo_salida_final']) else "No generado"
            filas_resumen.append(f"""
            <tr>
                <td style="padding: 12px; border-bottom: 1px solid #f0f0f0;"><strong>{r['regional']}</strong></td>
                <td style="padding: 12px; border-bottom: 1px solid #f0f0f0;">{estado_badge}</td>
                <td style="padding: 12px; border-bottom: 1px solid #f0f0f0;">{r['periodo']}</td>
                <td style="padding: 12px; border-bottom: 1px solid #f0f0f0;">{r['incidencias_resumen']}</td>
                <td style="padding: 12px; border-bottom: 1px solid #f0f0f0;">{excel_status}</td>
            </tr>
            """)
        
        tabla_resumen_html = f"""
        <table class="resumen-tabla" width="100%" style="border-collapse: collapse; width: 100%; background: white; margin-bottom: 20px;">
            <thead>
                <tr>
                    <th style="background-color: #e2e8f0; color: #2d3748; padding: 12px; font-weight: 700; border-bottom: 2px solid #cbd5e0; text-align: left;">Regional</th>
                    <th style="background-color: #e2e8f0; color: #2d3748; padding: 12px; font-weight: 700; border-bottom: 2px solid #cbd5e0; text-align: left;">Estado</th>
                    <th style="background-color: #e2e8f0; color: #2d3748; padding: 12px; font-weight: 700; border-bottom: 2px solid #cbd5e0; text-align: left;">Periodo</th>
                    <th style="background-color: #e2e8f0; color: #2d3748; padding: 12px; font-weight: 700; border-bottom: 2px solid #cbd5e0; text-align: left;">Resumen de Incidencias</th>
                    <th style="background-color: #e2e8f0; color: #2d3748; padding: 12px; font-weight: 700; border-bottom: 2px solid #cbd5e0; text-align: left;">Archivo Excel</th>
                </tr>
            </thead>
            <tbody>
                {"".join(filas_resumen)}
            </tbody>
        </table>
        """

        # Construir secciones de detalle de cada regional
        secciones_detalle = []
        archivos_adjuntos = []
        
        for r in self.resultados_piloto_actual:
            regional = r['regional']
            if r['proceso_ok']:
                if r['archivo_salida_final'] and os.path.exists(r['archivo_salida_final']):
                    archivos_adjuntos.append(r['archivo_salida_final'])
                
                # Generar el cuerpo de correo individual para esta regional
                cuerpo_ind = self.construir_cuerpo_correo(r['reporte_incidentes'], regional)
                # Extraemos el contenido dentro de <div class="content">...</div>
                content_start = cuerpo_ind.find('<div class="content">')
                content_end = cuerpo_ind.rfind('</div>\n                <div class="footer">')
                
                if content_start != -1 and content_end != -1:
                    inner_content = cuerpo_ind[content_start + len('<div class="content">'):content_end]
                else:
                    inner_content = cuerpo_ind
                
                # Quitar el saludo individual y la nota de archivo adjunto si están presentes
                saludo_idx = inner_content.find('<strong>Buen')
                if saludo_idx != -1:
                    end_p_saludo = inner_content.find('</p>', saludo_idx)
                    if end_p_saludo != -1:
                        inner_content = inner_content[end_p_saludo + 4:]
                
                adjunto_idx = inner_content.find('<div class="attachment-note">')
                if adjunto_idx != -1:
                    end_div_adj = inner_content.find('</div>', adjunto_idx)
                    if end_div_adj != -1:
                        inner_content = inner_content[end_div_adj + 6:]
                
                secciones_detalle.append(f"""
                <div class="regional-section" style="margin-top: 40px; border-top: 2px solid #e2e8f0; padding-top: 25px;">
                    <div class="regional-title" style="font-size: 18px; font-weight: 700; color: #2c3e50; margin-bottom: 15px; padding-bottom: 5px; border-bottom: 2px solid #cbd5e0;">📍 Detalles Regional: {regional}</div>
                    {inner_content}
                </div>
                """)
            else:
                secciones_detalle.append(f"""
                <div class="regional-section" style="margin-top: 40px; border-top: 2px solid #e2e8f0; padding-top: 25px;">
                    <div class="regional-title" style="font-size: 18px; font-weight: 700; color: #2c3e50; margin-bottom: 15px; padding-bottom: 5px; border-bottom: 2px solid #cbd5e0;">📍 Detalles Regional: {regional}</div>
                    <div class="error-box" style="background-color: #fff5f5; border-left: 4px solid #e53e3e; border-radius: 8px; padding: 20px; color: #c53030; line-height: 1.8;">
                        <strong>⚠️ Fallo en la Ejecución:</strong><br>
                        Periodo: {r['periodo']}<br>
                        Error: {r['incidencias_resumen']}<br>
                        Por favor, revise los logs del sistema para más información detallada del problema.
                    </div>
                </div>
                """)
        
        saludo = "Buen día" if datetime.now().hour < 12 else "Buenas tardes"
        
        cuerpo_html = f"""
        <html>
        <head>
            <meta http-equiv="Content-Type" content="text/html; charset=utf-8">
            <meta name="viewport" content="width=device-width, initial-scale=1.0">
            <style type="text/css">
                body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, Arial, sans-serif; font-size: 14px; line-height: 1.6; color: #333; margin: 0; padding: 20px; background-color: #f5f7fa; }}
                .email-container {{ max-width: 900px; margin: 0 auto; background: white; border-radius: 12px; overflow: hidden; box-shadow: 0 8px 16px rgba(0,0,0,0.1); }}
                .header {{ background-color: #4a5568; color: white; padding: 30px; text-align: left; }}
                .content {{ padding: 30px; }}
                table {{ border-collapse: collapse; width: 100%; background: white; margin-bottom: 20px; }}
                th, td {{ text-align: left; padding: 12px 15px; border-bottom: 1px solid #f0f0f0; }}
                th {{ background-color: #f8f9fa; font-weight: 600; color: #495057; font-size: 13px; text-transform: uppercase; }}
                .resumen-tabla th {{ background-color: #e2e8f0; color: #2d3748; font-weight: 700; }}
                .regional-section {{ margin-top: 40px; border-top: 2px solid #e2e8f0; padding-top: 25px; }}
                .regional-title {{ font-size: 18px; font-weight: 700; color: #2c3e50; margin-bottom: 15px; padding-bottom: 5px; border-bottom: 2px solid #cbd5e0; }}
                .error-box {{ background-color: #fff5f5; border-left: 4px solid #e53e3e; border-radius: 8px; padding: 20px; color: #c53030; line-height: 1.8; }}
                .attachment-note {{ background-color: #ebf8ff; border-left: 4px solid #3182ce; padding: 12px 20px; border-radius: 8px; margin: 20px 0; font-size: 14px; color: #2b6cb0; }}
                .info-box {{ background-color: #e3f2fd; border-left: 4px solid #2196f3; border-radius: 8px; padding: 20px; margin: 20px 0; }}
                .porc-sinc-container {{ background-color: #f5f7fa; border-radius: 12px; padding: 25px; margin: 25px 0; box-shadow: 0 4px 12px rgba(0,0,0,0.08); }}
                .porc-sinc-title {{ font-size: 20px; font-weight: 700; color: #2c3e50; margin-bottom: 20px; text-align: center; padding-bottom: 15px; border-bottom: 3px solid #667eea; }}
                .stat-card {{ background-color: white; border-radius: 10px; padding: 20px; margin: 10px 0; box-shadow: 0 3px 8px rgba(0,0,0,0.08); border-top: 4px solid #667eea; text-align: center; }}
                .stat-value {{ font-size: 32px; font-weight: 700; color: #2c3e50; margin: 5px 0; }}
                .stat-label {{ font-size: 12px; color: #7f8c8d; text-transform: uppercase; letter-spacing: 0.5px; font-weight: 600; }}
                .exec-today-pct {{ font-size: 13px; color: #7f8c8d; background: #f0f2f5; display: inline-block; padding: 2px 10px; border-radius: 12px; margin-top: 8px; font-weight: 600; border: 1px solid #e1e4e8; }}
                .info-note {{ background-color: #e8f5e9; border-left: 4px solid #4caf50; padding: 15px 20px; border-radius: 8px; margin: 20px 0; font-size: 13px; color: #1b5e20; }}
                .footer {{ margin-top: 40px; padding: 25px 30px; background-color: #f8f9fa; border-top: 3px solid #4a5568; }}
            </style>
        </head>
        <body>
            <div class="email-container">
                <div class="header">
                    <h2 style="margin: 0; font-size: 24px; color: white; font-weight: 700;">📊 Reporte Consolidado de Sincronización</h2>
                    <p style="margin: 8px 0 0 0; font-size: 14px; color: white;">Piloto Automático | ADA</p>
                </div>
                
                <div class="content">
                    <p style="font-size: 16px; margin-bottom: 20px; color: #2c3e50;"><strong>{saludo},</strong></p>
                    <p>Se ha completado el ciclo de ejecución programado del piloto automático. A continuación se presenta el resumen consolidado de las regionales analizadas.</p>
                    
                    {"".join([f'<div class="attachment-note">📎 <strong>Archivos adjuntos:</strong> Se incluyen {len(archivos_adjuntos)} reportes detallados en formato Excel.</div>' if archivos_adjuntos else ''])}
                    
                    <div class="info-box">
                        <h3 style="margin: 0 0 10px 0; font-size: 16px; color: #1565c0;">📋 Información General</h3>
                        <ul style="margin: 0; padding-left: 20px;">
                            <li><strong>Fecha de Ejecución:</strong> {fecha_actual}</li>
                            <li><strong>Regionales Procesadas:</strong> {", ".join(regiones_ejecutadas)}</li>
                        </ul>
                    </div>
                    
                    <h3 style="margin-top: 30px; color: #2c3e50;">📈 Resumen de Ejecución</h3>
                    {tabla_resumen_html}
                    
                    {"".join(secciones_detalle)}
                </div>
                
                <div class="footer">
                    <p style="margin: 0; font-size: 12px; color: #7f8c8d; text-align: center;">Este es un correo automático generado por el Sistema ADA.</p>
                </div>
            </div>
        </body>
        </html>
        """

        # Enviar correo de forma asíncrona
        threading.Thread(
            target=enviar_correo_notificacion,
            args=(asunto, cuerpo_html, destinatarios, archivos_adjuntos),
            daemon=True
        ).start()

    def update_incidents_table(self, incidents_df):
        for item in self.incidents_tree.get_children(): self.incidents_tree.delete(item)
        if incidents_df is not None and not incidents_df.empty:
            for i, row in incidents_df.iterrows():
                tag = 'oddrow' if i % 2 == 0 else 'evenrow'
                self.incidents_tree.insert('', END, values=list(row), tags=(tag,))
    
    def save_settings(self):
        settings = {
            'last_folder': self.carpeta_entrada.get(),
            'start_date': self.date_start.entry.get(),
            'piloto_activo': self.piloto_activo,
            'piloto_dias': {dia: var.get() for dia, var in self.vars_dias.items()},
            'proxima_ejecucion_ts': self.proxima_ejecucion_dt.timestamp() if self.piloto_activo and self.proxima_ejecucion_dt else None,
            'execution_history': self.execution_history,
            'email_recipients': self.recipients_list,
            'selected_regionales': {regional: var.get() for regional, var in self.regional_vars.items()},  # Guardar regionales
            'supabase_url': getattr(self, 'supabase_url', ''),
            'supabase_key': getattr(self, 'supabase_key', '')
        }
        try:
            with open(SETTINGS_FILE, 'w') as f:
                json.dump(settings, f, indent=4)
        except Exception as e:
            logging.error(f"No se pudo guardar la configuración: {e}")


    def load_settings(self):
        try:
            if os.path.exists(SETTINGS_FILE):
                with open(SETTINGS_FILE, 'r') as f:
                    settings = json.load(f)

                self.carpeta_entrada.set(settings.get('last_folder', ''))

                start_date = settings.get('start_date')
                if start_date:
                    self.date_start.entry.delete(0, END)
                    self.date_start.entry.insert(0, start_date)
                self.date_end.entry.delete(0, END)
                self.date_end.entry.insert(0, datetime.now().strftime('%d/%m/%Y'))

                self.execution_history = settings.get('execution_history', [])
                self.update_history_table()

                # Cargar destinatarios
                self.recipients_list = settings.get('email_recipients', [])
                for recipient in self.recipients_list:
                    if 'activo' not in recipient:
                        recipient['activo'] = True
                    activo_str = "☑" if recipient.get('activo', True) else "☐"
                    self.recipients_tree.insert('', END, values=(recipient.get('nombre'), recipient.get('correo'), recipient.get('cargo'), recipient.get('regional'), activo_str))

                # Restaurar estado de regionales seleccionadas
                regionales_guardadas = settings.get('selected_regionales', {})
                for regional, var in self.regional_vars.items():
                    if regionales_guardadas.get(regional):
                        var.set(True)

                self.supabase_url = settings.get('supabase_url', 'https://jcozaaifpfukqlypfuqq.supabase.co')
                self.supabase_key = settings.get('supabase_key', 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6Impjb3phYWlmcGZ1a3FseXBmdXFxIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4MDY4Mjg4MiwiZXhwIjoyMDk2MjU4ODgyfQ.80jkincEydxGbuhqKYEjJBOf1LJ3TkhthVgl0sPah8k')

                if settings.get('piloto_activo'):
                    self.piloto_switch.invoke()
                    dias_guardados = settings.get('piloto_dias', {})
                    for dia_key, var in self.vars_dias.items():
                        if dias_guardados.get(dia_key):
                            var.set(True)

                    ts = settings.get('proxima_ejecucion_ts')
                    if ts:
                        proxima_dt = datetime.fromtimestamp(ts)
                        if proxima_dt > datetime.now():
                            self.iniciar_piloto(from_load=True, proxima_ejecucion_guardada=proxima_dt)
                        else:
                            fecha_str = proxima_dt.strftime('%d/%m/%Y %H:%M:%S')
                            confirm = Messagebox.show_question(
                                f'¿Desea ejecutar el proceso pendiente del piloto automático ahora?\n(Programado originalmente para el {fecha_str})',
                                'Ejecución Pendiente',
                                parent=self
                            )
                            if confirm == 'Yes':
                                self.iniciar_piloto(from_load=False)
                            else:
                                self.iniciar_piloto(from_load=True, proxima_ejecucion_guardada=None)

        except Exception as e:
            logging.error(f"No se pudo cargar la configuración: {e}")
            Messagebox.show_error(f"Error al cargar configuración: {e}", "Error de Configuración", parent=self)


class LoginWindow(ttk.Toplevel):
    def __init__(self, parent):
        super().__init__(master=parent, title="Iniciar Sesión")
        self.parent = parent
        self.geometry("400x480")
        self.resizable(False, False)
        
        # Centrar ventana
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        x = (screen_width - 400) // 2
        y = (screen_height - 480) // 2
        self.geometry(f"+{x}+{y}")
        
        try: self.iconbitmap(resource_path('VicTor.ico'))
        except: pass

        self.protocol("WM_DELETE_WINDOW", self.on_close)

        style = ttk.Style()
        style.configure('Login.TFrame') # Usar fondo por defecto del tema

        main_frame = ttk.Frame(self, padding=30, style='Login.TFrame')
        main_frame.pack(fill=BOTH, expand=True)

        # Título / Logo
        ttk.Label(main_frame, text="Bienvenido", font=("Helvetica", 24, "bold")).pack(pady=(20, 10))
        ttk.Label(main_frame, text="Analizador BD by VicTor", font=("Helvetica", 10), bootstyle="light").pack(pady=(0, 30))

        # Usuario
        ttk.Label(main_frame, text="Usuario").pack(anchor=W)
        self.user_entry = ttk.Entry(main_frame, bootstyle="secondary")
        self.user_entry.pack(fill=X, pady=(5, 15))
        self.user_entry.insert(0, "vdominguez")  # Prellenar con usuario predeterminado

        # Contraseña
        ttk.Label(main_frame, text="Contraseña").pack(anchor=W)
        self.pass_entry = ttk.Entry(main_frame, show="*", bootstyle="secondary")
        self.pass_entry.pack(fill=X, pady=(5, 5))
        
        
        # Focus en contraseña
        self.pass_entry.bind("<Return>", self.validar_login)
        self.user_entry.bind("<Return>", lambda e: self.pass_entry.focus())

        # Link Recuperar
        lbl_recover = ttk.Label(main_frame, text="¿Olvidaste tu contraseña?", font=("Helvetica", 9, "underline"), bootstyle="info", cursor="hand2")
        lbl_recover.pack(anchor=E, pady=(0, 20))
        lbl_recover.bind("<Button-1>", self.recuperar_contrasena)

        # Botón
        self.btn_login = ttk.Button(main_frame, text="INICIAR SESIÓN", bootstyle="success", command=self.validar_login, width=20)
        self.btn_login.pack(pady=10, ipady=5)

        # Footer
        ttk.Label(main_frame, text="© 2025 Víctor Domínguez", font=("Helvetica", 8), bootstyle="light").pack(side=BOTTOM, pady=10)
        
        # Asegurar que el cursor esté en contraseña después de renderizar
        self.after(100, lambda: self.pass_entry.focus_set())

    def validar_login(self, event=None):
        usuario = self.user_entry.get().strip()
        contrasena = self.pass_entry.get().strip()

        if usuario == "vdominguez" and contrasena == "vicmat04":
            self.destroy()
            self.parent.deiconify() # Mostrar ventana principal
            self.parent.state('zoomed') # Maximizar ventana
            self.parent.load_settings()
        else:
            self.shake_window()
            Messagebox.show_error("Usuario o contraseña incorrectos.", "Error de Acceso", parent=self)
            self.pass_entry.delete(0, END)

    def on_close(self):
        self.destroy()
        self.parent.destroy() # Cerrar app si cierra login

    def shake_window(self):
        x = self.winfo_x()
        y = self.winfo_y()
        for i in range(0, 5):
            self.geometry(f"+{x+5}+{y}")
            self.update()
            time.sleep(0.04)
            self.geometry(f"+{x-5}+{y}")
            self.update()
            time.sleep(0.04)
        self.geometry(f"+{x}+{y}")

    def recuperar_contrasena(self, event=None):
        confirm = Messagebox.show_question("¿Desea enviar la contraseña a su correo registrado?", "Recuperar Contraseña", parent=self)
        if confirm == 'Yes':
            self.btn_login.config(state=DISABLED, text="Enviando...")
            self.update()
            
            cuerpo = """
            <html><body>
                <h3>Recuperación de Contraseña</h3>
                <p>Usted ha solicitado recuperar su contraseña para el <b>Analizador BD by VicTor</b>.</p>
                <p>Sus credenciales son:</p>
                <ul>
                    <li><b>Usuario:</b> vdominguez</li>
                    <li><b>Contraseña:</b> vicmat04</li>
                </ul>
                <p>Por favor, mantenga esta información segura.</p>
            </body></html>
            """
            
            t = threading.Thread(target=self._enviar_correo_recuperacion, args=("Recuperación de Contraseña - Analizador BD", cuerpo))
            t.start()
    
    def _enviar_correo_recuperacion(self, asunto, cuerpo):
        destinatarios = ["vicmat04@gmail.com", "victorpty999@gmail.com"]
        exito = enviar_correo_notificacion(asunto, cuerpo, destinatarios)
        self.after(0, lambda: self._post_recuperacion(exito))

    def _post_recuperacion(self, exito):
        self.btn_login.config(state=NORMAL, text="INICIAR SESIÓN")
        if exito:
            Messagebox.show_info("La contraseña ha sido enviada a su correo.", "Correo Enviado", parent=self)
        else:
            Messagebox.show_error("No se pudo enviar el correo. Verifique su conexión.", "Error de Envío", parent=self)


if __name__ == "__main__":
    app = App()
    app.withdraw() # Ocultar ventana principal
    login_app = LoginWindow(app) # Pasar app como master
    app.mainloop()